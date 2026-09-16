"""Workbook inventory: builds a WorkbookAnalysis-shaped dict from a copy of
the source file using openpyxl (cells/formulas/names/styles) and zipfile
(raw OOXML package parts). Never executes VBA, links, or OLE content --
inspection only, per instructions/core.md.

The `workbooks[i]` / `features` / `risks` shapes here are this app's own
convention: schemas/workbook-analysis.schema.json declares those fields as
plain `type: object` / `type: array` with no nested schema, so there is
nothing more specific to conform to.

1.3.0 additions (all read-only, all from the same single cell pass):
  * Mind-style grid detection with '#Title /Flag' parsing (app/grids.py)
  * array formulas (legacy CSE and dynamic-array) and data-table formulas,
    which openpyxl hands back as objects rather than '=...' strings and
    were previously invisible to every formula check
  * per-sheet style statistics (number-format categories, theme colours,
    empty-but-filled cells), outlines/hidden rows+columns, data validations,
    conditional formatting, sheet protection, tab colour
  * docProps/app.xml (which application last saved the file) and the size of
    the VBA project part
  * a process-wide cache of the opened workbook so validators that need raw
    cell values (resize destinations, MM_READTABLE ranges) don't re-parse a
    multi-MB file.
"""
from __future__ import annotations

import hashlib
import re
import shutil
import zipfile
from pathlib import Path
from typing import Any, Callable, Iterable

import openpyxl
from openpyxl.reader.excel import ExcelReader
from openpyxl.styles.numbers import is_date_format
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from .formula_utils import called_functions, find_calls, range_length, storage_prefixes
from .grids import detect_grids, is_occupied, json_safe

STEP_LABELS_MARKER = "steplabels"
MAX_STYLE_SAMPLES = 25
MAX_STYLE_CELL_REFS = 5000  # coordinates kept per sheet for app/prep.py actions

_WORKBOOK_CACHE: dict[str, Any] = {}

# A progress callback: progress(stage, message, fraction_within_stage_or_None, **facts).
# Stages, in order: copy, load, inventory, names, rules, plan (the web server
# turns them into one status the front-end polls). Any exception a callback
# raises is the caller's problem -- the scanner never swallows it.
ProgressFn = Callable[..., None]


def _notify(progress: ProgressFn | None, stage: str, message: str, fraction: float | None = None, **facts: Any) -> None:
    if progress is not None:
        progress(stage, message, fraction, **facts)


class _SelectiveReader(ExcelReader):
    """openpyxl's reader, except that the sheets in `ignore` are never parsed:
    they come back as empty placeholder sheets at their original position
    (so sheet indexes, sheet-scoped defined names and the active-sheet index
    stay right) with their real name and visibility. Everything else --
    strings, styles, names, VBA, the other sheets -- is read exactly as
    openpyxl would. Also the hook for per-sheet load progress: openpyxl
    parses one sheet at a time and the workbook part has told us how many."""

    def __init__(self, fn, ignore: set[str], progress: ProgressFn | None, **kwargs):
        super().__init__(fn, **kwargs)
        self._ignore = ignore
        self._progress = progress

    def read_worksheets(self):
        parser = self.parser
        original = parser.find_sheets
        skipped: list[tuple[int, Any]] = []

        def filtered():
            sheets = list(original())
            total = len(sheets)
            for i, (sheet, rel) in enumerate(sheets):
                if sheet.name in self._ignore:
                    skipped.append((i, sheet))
                    _notify(self._progress, "load", f"Skipping sheet {i + 1}/{total}: {sheet.name}", i / total, sheet=sheet.name, skipped=True)
                    continue
                _notify(self._progress, "load", f"Reading sheet {i + 1}/{total}: {sheet.name}", i / total, sheet=sheet.name)
                yield sheet, rel

        parser.find_sheets = filtered
        try:
            super().read_worksheets()
        finally:
            parser.find_sheets = original
        for i, sheet in skipped:  # ascending index order keeps every position right
            ws = self.wb.create_sheet(sheet.name, index=i)
            ws.sheet_state = sheet.state
            ws.mind_ready_ignored = True


def load_workbook_selective(path: str | Path, ignore_sheets: Iterable[str] = (), data_only: bool = False, keep_vba: bool = False, progress: ProgressFn | None = None):
    """openpyxl.load_workbook that does not parse the sheets in `ignore_sheets`
    (they are present, but empty). With nothing to ignore and no progress
    hook it is exactly openpyxl.load_workbook."""
    ignore = {str(n) for n in ignore_sheets or ()}
    if not ignore and progress is None:
        return openpyxl.load_workbook(str(path), data_only=data_only, keep_vba=keep_vba)
    reader = _SelectiveReader(str(path), ignore, progress, read_only=False, keep_vba=keep_vba, data_only=data_only, keep_links=True, rich_text=False)
    reader.read()
    return reader.wb


def ignored_sheet_names(analysis: dict[str, Any], workbook_index: int = 0) -> list[str]:
    """The sheets the user asked the scan to skip (none for an ordinary analysis)."""
    try:
        return [s["name"] for s in analysis["workbooks"][workbook_index].get("ignored_sheets", [])]
    except (KeyError, IndexError, TypeError):
        return []


def sha256_of(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def has_vba_project(path: Path) -> bool:
    """Cheap zipfile-only check for a VBA project part -- works even on
    .xlsb (also an OPC/zip container, confirmed empirically: 'xl/vbaProject.bin'
    is present the same way it is in .xlsm), which openpyxl can't open at all.
    Used to decide whether an auto-converted file should target .xlsm
    (preserve macros) or .xlsx (no macros to preserve)."""
    try:
        with zipfile.ZipFile(path) as zf:
            return "xl/vbaProject.bin" in zf.namelist()
    except zipfile.BadZipFile:
        return False


def make_immutable_copy(source_path: Path, work_dir: Path) -> tuple[Path, str]:
    """Hash the source, then copy it -- never open or modify the original.
    Non-negotiable rule #1 (SKILL.md): never modify the source workbook.
    """
    source_path = Path(source_path)
    if not source_path.is_file():
        raise FileNotFoundError(source_path)
    source_sha256 = sha256_of(source_path)
    work_dir.mkdir(parents=True, exist_ok=True)
    copy_path = work_dir / source_path.name
    # The source may itself be an earlier output living in `work_dir` (the UI
    # re-applies changes to the file it just produced): never copy a file onto
    # itself, and never overwrite a previous output -- use a fresh versioned
    # sub-folder instead.
    if copy_path.exists():
        n = 2
        while True:
            candidate = work_dir / f"v{n}" / source_path.name
            if not candidate.exists():
                candidate.parent.mkdir(parents=True, exist_ok=True)
                copy_path = candidate
                break
            n += 1
    shutil.copy2(source_path, copy_path)
    return copy_path, source_sha256


def open_cached(analysis: dict[str, Any], workbook_index: int = 0):
    """The openpyxl workbook behind an analysis (the immutable copy), kept
    open in a process-wide cache so validators can read arbitrary cells
    without re-parsing. Read-only use only."""
    copy_path = analysis["workbooks"][workbook_index]["copy_path"]
    wb = _WORKBOOK_CACHE.get(copy_path)
    if wb is None:
        wb = load_workbook_selective(copy_path, ignored_sheet_names(analysis, workbook_index), data_only=False)
        _WORKBOOK_CACHE[copy_path] = wb
    return wb


def cell_value(analysis: dict[str, Any], sheet: str, row: int, col: int) -> Any:
    wb = open_cached(analysis)
    if sheet not in wb.sheetnames:
        return None
    return json_safe(wb[sheet].cell(row, col).value)


_VALUES_CACHE: dict[str, Any] = {}


def _values_workbook(path: str, ignore_sheets: Iterable[str] = ()):
    """A data_only (cached values) view of a workbook, cached per path -- the
    values Excel last stored, i.e. what the user sees in the cells. Sheets the
    analysis skipped are skipped here too (they would cost the same again)."""
    ignore = tuple(sorted({str(n) for n in ignore_sheets or ()}))
    key = f"{path}|{'|'.join(ignore)}" if ignore else path
    wb = _VALUES_CACHE.get(key)
    if wb is None:
        wb = load_workbook_selective(path, ignore, data_only=True)
        _VALUES_CACHE[key] = wb
    return wb


def values_cell(analysis: dict[str, Any], sheet: str, row: int, col: int) -> Any:
    """The value Excel last stored for a cell (what a formula computed), or
    None when the workbook carries no cached value for it."""
    copy_path = analysis["workbooks"][0]["copy_path"]
    try:
        wb = _values_workbook(str(copy_path))
    except Exception:
        return None
    if sheet not in wb.sheetnames:
        return None
    try:
        return json_safe(wb[sheet].cell(row, col).value)
    except Exception:
        return None


def cell_window(analysis: dict[str, Any], sheet: str, cell: str, rows: int = 3, cols: int = 3, values_path: str | Path | None = None) -> dict[str, Any]:
    """What is going on in the workbook around `cell`: a small grid of the
    cells' contents (values) with their formulas, so a finding or a
    recalculation error can be understood in place. `values_path` may point
    at a fresher copy (e.g. the last recalculation) for the values."""
    from openpyxl.utils import get_column_letter

    from .formula_utils import parse_ref, ref_text

    wb_f = open_cached(analysis)
    if sheet not in wb_f.sheetnames:
        return {"error": f"sheet '{sheet}' not found", "sheet": sheet, "focus": cell, "columns": [], "rows": []}
    ref = parse_ref(str(cell).replace("$", "").upper())
    if not ref or ref["whole_column"] or ref["whole_row"]:
        return {"error": f"'{cell}' is not a cell or range", "sheet": sheet, "focus": cell, "columns": [], "rows": []}
    ws_f = wb_f[sheet]
    copy_path = analysis["workbooks"][0]["copy_path"]
    values_from = "analysis copy"
    ws_v = None
    for candidate, label in ((values_path, "recalculation"), (copy_path, "analysis copy")):
        if candidate and Path(candidate).is_file():
            try:
                wb_v = _values_workbook(str(candidate), ignored_sheet_names(analysis))
                if sheet in wb_v.sheetnames:
                    ws_v, values_from = wb_v[sheet], label
                    break
            except Exception:
                continue
    rows = max(0, min(int(rows), 10))
    cols = max(0, min(int(cols), 6))
    top, left = max(1, ref["r1"] - rows), max(1, ref["c1"] - cols)
    bottom = min(ref["r2"] + rows, max(ws_f.max_row, ref["r2"]))
    right = min(ref["c2"] + cols, max(ws_f.max_column, ref["c2"]))
    # multi-cell array formulas live in their anchor cell only; members show the anchor's formula
    arrays: list[tuple[dict, str]] = []
    for entry in analysis["workbooks"][0].get("formulas", []):
        if entry.get("sheet") == sheet and entry.get("array_ref") and ":" in str(entry["array_ref"]):
            r = parse_ref(str(entry["array_ref"]).replace("$", "").upper())
            if r:
                arrays.append((r, str(entry.get("formula") or "")))
    out_rows = []
    for r in range(top, bottom + 1):
        cells = []
        for c in range(left, right + 1):
            raw = ws_f.cell(r, c).value
            formula = None
            array_ref = None
            if isinstance(raw, str) and raw.startswith("="):
                formula = raw
            elif isinstance(raw, ArrayFormula):
                text = raw.text if str(raw.text).startswith("=") else "=" + str(raw.text)
                formula, array_ref = "{" + text + "}", str(raw.ref)
            elif isinstance(raw, DataTableFormula):
                formula = "=TABLE()"
            if formula is None:
                for ar, text in arrays:
                    if ar["r1"] <= r <= ar["r2"] and ar["c1"] <= c <= ar["c2"]:
                        formula, array_ref = "{" + text + "}", ar["ref"]
                        break
            value = json_safe(ws_v.cell(r, c).value) if ws_v is not None else (None if formula else json_safe(raw))
            cells.append({
                "ref": ref_text(c, r),
                "value": value,
                "formula": formula,
                "array": array_ref,
                "error": isinstance(value, str) and value.startswith("#"),
                "focus": ref["r1"] <= r <= ref["r2"] and ref["c1"] <= c <= ref["c2"],
            })
        out_rows.append({"row": r, "cells": cells})
    return {
        "sheet": sheet,
        "focus": ref["ref"],
        "range": f"{ref_text(left, top)}:{ref_text(right, bottom)}",
        "columns": [get_column_letter(c) for c in range(left, right + 1)],
        "rows": out_rows,
        "values_from": values_from,
    }


# --- number-format classification (FMT-001) ---------------------------------
_NUMERIC_FMT_RE = re.compile(r"^[#0,]*(\.[#0]+)?(;-?[#0,]*(\.[#0]+)?)?(;[^;]*)?$")


def classify_number_format(fmt: str | None, cell) -> str:
    """Map an Excel number format onto the KB's list of preserved cell
    formats ("Excel specificities kept in Mind": Standard, Number, Text,
    Boolean, Date) or a named category outside that list."""
    if cell.data_type == "b":
        return "boolean"
    if not fmt or fmt == "General":
        return "standard"
    if fmt == "@":
        return "text"
    if is_date_format(fmt):
        return "date"
    if "%" in fmt:
        return "percentage"
    if "E+" in fmt.upper() or "E-" in fmt.upper():
        return "scientific"
    if "?/" in fmt:
        return "fraction"
    if "_(" in fmt or "* " in fmt:
        return "accounting"
    if "$" in fmt or "€" in fmt or "£" in fmt or "[$" in fmt:
        return "currency"
    if _NUMERIC_FMT_RE.match(fmt.replace('"', "")):
        return "number"
    return "custom"


PRESERVED_FORMAT_CATEGORIES = {"standard", "number", "text", "boolean", "date"}


def _color_is_theme(color) -> bool:
    return color is not None and getattr(color, "type", None) == "theme"


def _app_properties(names_in_package: set[str], zf: zipfile.ZipFile) -> dict[str, Any]:
    props: dict[str, Any] = {"application": None, "app_version": None}
    if "docProps/app.xml" in names_in_package:
        xml = zf.read("docProps/app.xml").decode("utf-8", errors="replace")
        m = re.search(r"<Application>(.*?)</Application>", xml)
        if m:
            props["application"] = m.group(1).strip()
        m = re.search(r"<AppVersion>(.*?)</AppVersion>", xml)
        if m:
            props["app_version"] = m.group(1).strip()
    return props


def _sheet_inventory(ws) -> tuple[dict[str, Any], list[dict[str, Any]], dict[str, int]]:
    """One pass over the sheet's cells: occupancy map for grid detection,
    formulas, styles, comments, hyperlinks, locked cells."""
    merged = [str(r) for r in ws.merged_cells.ranges]
    sheet_protected = bool(getattr(ws.protection, "sheet", False))
    locked_cells = 0
    comments = []
    hyperlinks = []
    formulas: list[dict[str, Any]] = []
    function_usage: dict[str, int] = {}
    occupied: dict[tuple[int, int], Any] = {}
    theme_color_cells = 0
    theme_color_cell_refs: list[str] = []
    empty_styled_cells = 0
    empty_styled_cell_refs: list[str] = []
    format_categories: dict[str, int] = {}
    format_samples: list[dict[str, Any]] = []
    array_formula_count = 0
    prefix_counts: dict[str, int] = {}

    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            occupied_here = is_occupied(value)
            if occupied_here:
                occupied[(cell.row, cell.column)] = value

            formula_text: str | None = None
            entry: dict[str, Any] | None = None
            if isinstance(value, str) and value.startswith("="):
                formula_text = value
                entry = {"sheet": ws.title, "cell": cell.coordinate, "formula": value}
            elif isinstance(value, ArrayFormula):
                formula_text = value.text if value.text.startswith("=") else "=" + value.text
                entry = {"sheet": ws.title, "cell": cell.coordinate, "formula": formula_text, "array_ref": value.ref}
                array_formula_count += 1
            elif isinstance(value, DataTableFormula):
                entry = {"sheet": ws.title, "cell": cell.coordinate, "formula": "=TABLE()", "data_table": True}
                formula_text = None
            if entry is not None:
                if formula_text:
                    for fn in called_functions(formula_text):
                        function_usage[fn] = function_usage.get(fn, 0) + 1
                    prefixes = storage_prefixes(formula_text)
                    if prefixes:
                        entry["storage_prefixes"] = sorted(prefixes)
                        for p in prefixes:
                            prefix_counts[p] = prefix_counts.get(p, 0) + 1
                formulas.append(entry)

            # Excel locks every cell by default (an explicit lock leaves no style
            # trace), and locking only bites on a protected sheet (KB): count
            # occupied cells on protected sheets that are not explicitly unlocked.
            if occupied_here and sheet_protected and (cell.protection is None or cell.protection.locked):
                locked_cells += 1
            if cell.has_style:
                if occupied_here:
                    font_color = cell.font.color if cell.font is not None else None
                    fill = cell.fill
                    fill_theme = fill is not None and fill.fill_type is not None and _color_is_theme(fill.fgColor)
                    if _color_is_theme(font_color) or fill_theme:
                        theme_color_cells += 1
                        if len(theme_color_cell_refs) < MAX_STYLE_CELL_REFS:
                            theme_color_cell_refs.append(cell.coordinate)
                    category = classify_number_format(cell.number_format, cell)
                    format_categories[category] = format_categories.get(category, 0) + 1
                    if category not in PRESERVED_FORMAT_CATEGORIES and len(format_samples) < MAX_STYLE_SAMPLES:
                        format_samples.append({"cell": cell.coordinate, "number_format": cell.number_format, "category": category})
                else:
                    fill = cell.fill
                    if fill is not None and fill.fill_type == "solid":
                        empty_styled_cells += 1
                        if len(empty_styled_cell_refs) < MAX_STYLE_CELL_REFS:
                            empty_styled_cell_refs.append(cell.coordinate)
            elif occupied_here:
                category = "boolean" if cell.data_type == "b" else "standard"
                format_categories[category] = format_categories.get(category, 0) + 1

            if cell.comment is not None:
                comments.append({"cell": cell.coordinate, "author": cell.comment.author, "text": cell.comment.text})
            if cell.hyperlink is not None:
                hyperlinks.append({"cell": cell.coordinate, "target": cell.hyperlink.target})

    grids, standalone = detect_grids(ws.title, occupied)

    row_groups = hidden_rows = max_row_level = 0
    for dim in ws.row_dimensions.values():
        level = getattr(dim, "outline_level", 0) or 0
        if level:
            row_groups += 1
            max_row_level = max(max_row_level, level)
        if getattr(dim, "hidden", False):
            hidden_rows += 1
    col_groups = hidden_cols = max_col_level = 0
    for dim in ws.column_dimensions.values():
        level = getattr(dim, "outline_level", 0) or 0
        if level:
            col_groups += 1
            max_col_level = max(max_col_level, level)
        if getattr(dim, "hidden", False):
            hidden_cols += 1

    validations = []
    try:
        for dv in ws.data_validations.dataValidation:
            validations.append({"sqref": str(dv.sqref), "type": dv.type, "formula1": dv.formula1})
    except Exception:
        pass
    cf_rules = 0
    try:
        for cf in ws.conditional_formatting:
            cf_rules += len(cf.rules)
    except Exception:
        pass

    state = ws.sheet_state  # 'visible' | 'hidden' | 'veryHidden'
    tab_color = None
    try:
        if ws.sheet_properties.tabColor is not None:
            rgb = ws.sheet_properties.tabColor.rgb  # may be an openpyxl RGB descriptor, not a str
            tab_color = rgb if isinstance(rgb, str) else None
    except Exception:
        pass

    sheet_entry = {
        "name": ws.title,
        "state": state,
        "dimensions": ws.dimensions,
        "tab_color": tab_color,
        "protected": sheet_protected,
        "merged_cells": merged,
        "locked_cell_count": locked_cells,
        "comments": comments,
        "hyperlinks": hyperlinks,
        "grids": grids,
        "standalone_text_cells": standalone,
        "occupied_cell_count": len(occupied),
        "formula_count": len(formulas),
        "array_formula_count": array_formula_count,
        "outline": {
            "row_groups": row_groups,
            "column_groups": col_groups,
            "max_row_level": max_row_level,
            "max_column_level": max_col_level,
            "hidden_rows": hidden_rows,
            "hidden_columns": hidden_cols,
        },
        "data_validations": validations,
        "conditional_format_rule_count": cf_rules,
        "style_stats": {
            "theme_color_cells": theme_color_cells,
            "theme_color_cell_refs": theme_color_cell_refs,
            "empty_styled_cells": empty_styled_cells,
            "empty_styled_cell_refs": empty_styled_cell_refs,
            "number_format_categories": format_categories,
            "unsupported_format_samples": format_samples,
        },
        "formula_storage_prefixes": prefix_counts,
    }
    return sheet_entry, formulas, function_usage


def build_analysis(
    source_path: Path,
    work_dir: Path,
    analysis_id: str,
    ignore_sheets: Iterable[str] | None = None,
    progress: ProgressFn | None = None,
) -> dict[str, Any]:
    """The inventory of one workbook. `ignore_sheets` names sheets the user
    chose not to scan (a large-workbook option): they are neither parsed nor
    inventoried, appear under `workbooks[0].ignored_sheets` (name + state)
    and nowhere else, and are counted in `features.ignored_sheet_count`.
    `progress` receives the stages (see ProgressFn)."""
    ignore = {str(n) for n in ignore_sheets or ()}
    _notify(progress, "copy", "Copying the workbook (the original is never modified)", None)
    copy_path, source_sha256 = make_immutable_copy(source_path, work_dir)
    suffix = copy_path.suffix.lower().lstrip(".")

    _notify(progress, "load", "Opening the workbook", 0.0)
    wb = load_workbook_selective(copy_path, ignore, data_only=False, progress=progress)

    with zipfile.ZipFile(copy_path) as zf:
        names_in_package = set(zf.namelist())
        app_props = _app_properties(names_in_package, zf)
        vba_size = zf.getinfo("xl/vbaProject.bin").file_size if "xl/vbaProject.bin" in names_in_package else 0
        package_part_sizes = {i.filename: i.file_size for i in zf.infolist()}
    has_vba = "xl/vbaProject.bin" in names_in_package or suffix in ("xlsm", "xlsb")

    sheets: list[dict[str, Any]] = []
    ignored_sheets: list[dict[str, Any]] = []
    formulas: list[dict[str, Any]] = []
    function_usage: dict[str, int] = {}
    hidden_sheet_count = 0
    prefix_counts: dict[str, int] = {}

    all_sheets = list(wb.worksheets)
    for i, ws in enumerate(all_sheets):
        if ws.title in ignore:
            ignored_sheets.append({"name": ws.title, "state": ws.sheet_state, "index": i})
            continue
        _notify(progress, "inventory", f"Scanning sheet {i + 1}/{len(all_sheets)}: {ws.title}", i / max(1, len(all_sheets)), sheet=ws.title)
        sheet_entry, sheet_formulas, sheet_usage = _sheet_inventory(ws)
        sheets.append(sheet_entry)
        formulas.extend(sheet_formulas)
        for fn, n in sheet_usage.items():
            function_usage[fn] = function_usage.get(fn, 0) + n
        for p, n in sheet_entry["formula_storage_prefixes"].items():
            prefix_counts[p] = prefix_counts.get(p, 0) + n
        if sheet_entry["state"] != "visible":
            hidden_sheet_count += 1

    _notify(progress, "names", "Reading defined names and package parts", None)
    defined_names = []
    if hasattr(wb.defined_names, "values"):
        for dn in wb.defined_names.values():
            defined_names.append({"name": dn.name, "value": dn.attr_text, "local_sheet_id": dn.localSheetId})
    for idx, ws in enumerate(wb.worksheets):
        local = getattr(ws, "defined_names", None)
        if local and hasattr(local, "values"):
            for dn in local.values():
                defined_names.append({"name": dn.name, "value": dn.attr_text, "local_sheet_id": idx})

    external_link_parts = [n for n in names_in_package if n.startswith("xl/externalLinks/")]
    mm_functions_used = {k: v for k, v in function_usage.items() if k.startswith("MM_")}
    native_function_usage = {k: v for k, v in function_usage.items() if not k.startswith("MM_")}

    grid_count = sum(len(s["grids"]) for s in sheets)
    flagged_grid_count = sum(1 for s in sheets for g in s["grids"] if g["flags"])
    special_parts = sorted(
        n for n in names_in_package if n.startswith(("xl/model/", "customXml/", "xl/pivotCache/", "xl/slicers/", "xl/ctrlProps/", "xl/activeX/", "xl/embeddings/"))
    )

    features = {
        "sheet_count": len(sheets),
        "ignored_sheet_count": len(ignored_sheets),
        "total_sheet_count": len(all_sheets),
        "hidden_sheet_count": hidden_sheet_count,
        "has_vba": has_vba,
        "vba_project_bytes": vba_size,
        "function_usage": function_usage,
        "mm_functions_used": mm_functions_used,
        "native_function_usage": native_function_usage,
        "external_link_part_count": len(external_link_parts),
        "defined_name_count": len(defined_names),
        "formula_count": len(formulas),
        "array_formula_count": sum(s["array_formula_count"] for s in sheets),
        "formula_storage_prefixes": prefix_counts,
        "grid_count": grid_count,
        "flagged_grid_count": flagged_grid_count,
        "application": app_props["application"],
        "app_version": app_props["app_version"],
        "package_part_count": len(names_in_package),
        "non_openpyxl_parts": special_parts,
    }

    risks: list[dict[str, Any]] = []
    if has_vba:
        risks.append({"type": "VBA_PRESENT", "message": "Workbook contains a VBA project."})
    if external_link_parts:
        risks.append({"type": "EXTERNAL_LINKS_PRESENT", "message": f"{len(external_link_parts)} external link package part(s) found."})
    if hidden_sheet_count:
        risks.append({"type": "HIDDEN_SHEETS_PRESENT", "message": f"{hidden_sheet_count} hidden or very-hidden sheet(s)."})
    if special_parts:
        risks.append(
            {
                "type": "PARTS_OPENPYXL_CANNOT_PRESERVE",
                "message": f"{len(special_parts)} package part(s) (data model / customXml / pivot caches / controls) that a pure-Python save would drop; outputs must be written by Excel.",
            }
        )
    if ignored_sheets:
        risks.append(
            {
                "type": "SHEETS_IGNORED",
                "message": f"{len(ignored_sheets)} sheet(s) skipped at the user's request and not checked by any rule: "
                + ", ".join(s["name"] for s in ignored_sheets),
            }
        )

    workbook_entry = {
        "path": str(source_path),
        "copy_path": str(copy_path),
        "file_name": Path(source_path).stem,
        "file_type": suffix,
        "source_sha256": source_sha256,
        "sheets": sheets,
        "ignored_sheets": ignored_sheets,
        "defined_names": defined_names,
        "formulas": formulas,
        "has_vba": has_vba,
        "external_link_parts": external_link_parts,
        "package_part_sizes": package_part_sizes,
    }

    _WORKBOOK_CACHE[str(copy_path)] = wb

    return {
        "schema_version": "1.0.0",
        "analysis_id": analysis_id,
        "source": {"path": str(source_path), "sha256": source_sha256, "copy_path": str(copy_path)},
        "workbooks": [workbook_entry],
        "features": features,
        "risks": risks,
    }


# --- MM-call iterators shared by the validators -----------------------------
def mm_calls(analysis: dict[str, Any], function_name: str) -> list[dict[str, Any]]:
    """Every call site of one MM_ function: {sheet, cell, formula, args, array_ref}."""
    out = []
    for wb_entry in analysis["workbooks"]:
        for f in wb_entry["formulas"]:
            for args in find_calls(f["formula"], function_name):
                out.append(
                    {
                        "sheet": f["sheet"],
                        "cell": f["cell"],
                        "formula": f["formula"],
                        "args": args,
                        "array_ref": f.get("array_ref"),
                    }
                )
    return out


def loop_definitions(analysis: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    """Every MM_LOOP("Name", Range, ...) call site, grouped by Name."""
    defs: dict[str, list[dict[str, Any]]] = {}
    for call in mm_calls(analysis, "MM_LOOP"):
        args = call["args"]
        if not args:
            continue
        name = args[0].strip().strip('"')
        range_arg = args[1].strip() if len(args) > 1 else None
        defs.setdefault(name, []).append(
            {
                "sheet": call["sheet"],
                "cell": call["cell"],
                "formula": call["formula"],
                "range_arg": range_arg,
                "range_length": range_length(range_arg) if range_arg else None,
                "literal_name": args[0].strip().startswith('"'),
            }
        )
    return defs


def loopinstance_definitions(analysis: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    defs: dict[str, list[dict[str, Any]]] = {}
    for call in mm_calls(analysis, "MM_LOOPINSTANCE"):
        args = call["args"]
        if not args:
            continue
        name = args[0].strip().strip('"')
        defs.setdefault(name, []).append({"sheet": call["sheet"], "cell": call["cell"], "formula": call["formula"], "arg_count": len(args)})
    return defs


def looplabels_definitions(analysis: dict[str, Any]) -> list[dict[str, Any]]:
    """Every MM_LOOPLABELS("Name", Range, ...) call site."""
    out = []
    for call in mm_calls(analysis, "MM_LOOPLABELS"):
        args = call["args"]
        if not args:
            continue
        name = args[0].strip().strip('"')
        range_arg = args[1].strip() if len(args) > 1 else None
        out.append(
            {
                "sheet": call["sheet"],
                "cell": call["cell"],
                "name": name,
                "range_arg": range_arg,
                "range_length": range_length(range_arg) if range_arg else None,
            }
        )
    return out


def known_loop_names(analysis: dict[str, Any]) -> set[str]:
    """Loop names a formula may legitimately reference: MM_LOOP / MM_LOOPINSTANCE
    definitions, iteration names (MM_ITERATIONS), the built-in SIM (simulations),
    BATCH (run batches) and Analysis (AOC/AOS) dimensions -- KB MM_RESULT,
    MM_DIMSIZE, MM_BATCHINDEX and AOC/AOS articles."""
    names = set(loop_definitions(analysis)) | set(loopinstance_definitions(analysis))
    for call in mm_calls(analysis, "MM_ITERATIONS"):
        if call["args"]:
            names.add(call["args"][0].strip().strip('"'))
    names |= {"SIM", "BATCH", "Analysis"}
    return names
