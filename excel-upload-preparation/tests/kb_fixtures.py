"""Synthetic 'Mind-style' model fixtures for the 1.3.0 KB-backed validators
(registered from tests/conftest.py via pytest_plugins). Same conventions as
conftest.py: built fresh with openpyxl into tmp_path, labelled
SYNTHETIC_TEST_FIXTURE in A1, no binary files checked in."""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

LABEL = "SYNTHETIC_TEST_FIXTURE"


def _put(ws, ref: str, value) -> None:
    ws[ref] = value


@pytest.fixture
def flagged_model_xlsx(tmp_path: Path) -> Path:
    """A small but complete model that follows the KB: titled, flagged grids
    separated by empty columns/rows. Every 1.3.0 validator that has something
    to check here should PASS (or WARNING where the KB only asks for review)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    _put(ws, "A1", LABEL)  # a standalone text cell (STR-002 reports it -- expected)

    # /Input + /Reorder grid with unique text headers
    _put(ws, "A3", "#Assumptions /Input /Reorder")
    for col, h in zip("ABC", ("Key", "Value", "Note")):
        _put(ws, f"{col}4", h)
    for i, (k, v) in enumerate((("alpha", 1), ("beta", 2), ("gamma", 3)), start=5):
        _put(ws, f"A{i}", k)
        _put(ws, f"B{i}", v)
        _put(ws, f"C{i}", "n")

    # /Parameters grid
    _put(ws, "E3", "#Params /Parameters")
    for col, h in zip("EFGH", ("Label", "Type", "PossibleValues", "Values")):
        _put(ws, f"{col}4", h)
    for i, row in enumerate((("Rate", "number", "0|1|0.1", 0.5), ("Mode", "dropdown", "A|B", "A"), ("Flag", "switch", "", "true")), start=5):
        for col, v in zip("EFGH", row):
            _put(ws, f"{col}{i}", v)

    # loop + result
    _put(ws, "J3", "#Scenarios")
    _put(ws, "J4", "Scenario")
    for i, v in enumerate((10, 20, 30), start=5):
        _put(ws, f"J{i}", v)
    _put(ws, "L3", "#Calc")
    _put(ws, "L4", "Loop")
    _put(ws, "L5", '=MM_LOOP("Scenario", J5:J7)')
    _put(ws, "L6", '=MM_RESULT(L5,"Scenario",1)')

    # resize group: reference grid with MM_SETSIZE + follower grid
    _put(ws, "N3", "#Sizes /Resize.grp")
    _put(ws, "N4", "N")
    _put(ws, "N5", 3)
    _put(ws, "P3", "#Table /Resize.grp")
    _put(ws, "P4", "Val")
    _put(ws, "P5", "=N5*2+MM_SETSIZE(N5,1)")  # P6, P7 left empty for the copies

    # group of grids
    _put(ws, "R3", "#G1 /Group.(Block).0.0")
    _put(ws, "R4", "h")
    _put(ws, "R5", 1)
    _put(ws, "T3", "#G2 /Group.(Block).0.1")
    _put(ws, "T4", "h")
    _put(ws, "T5", 2)

    # backup source/dest + button outside them
    _put(ws, "V3", "#Src /BackupSource.(b1)")
    _put(ws, "V4", "h")
    _put(ws, "V5", 1)
    _put(ws, "X3", "#Dst /BackupDest.(b1)")
    _put(ws, "X4", "h")
    _put(ws, "X5", 1)
    _put(ws, "Z3", "#Btn")
    _put(ws, "Z4", "h")
    _put(ws, "Z5", '=MM_BACKUPBUTTON("Run backup")')

    # /HideRows column (last column), #NbSimulations, /Translations
    _put(ws, "AB3", "#Hide")
    _put(ws, "AB4", "h")
    _put(ws, "AC4", "/HideRows")
    _put(ws, "AB5", 1)
    _put(ws, "AC5", True)
    _put(ws, "AE3", "#NbSimulations")
    _put(ws, "AE4", 1000)
    _put(ws, "AG3", "#Tr /Translations")
    _put(ws, "AG4", "en")
    _put(ws, "AH4", "fr")
    _put(ws, "AG5", "Hello")
    _put(ws, "AH5", "Bonjour")

    # MM_READTABLENAN over a headed table, plus MM_RANGE alone in its own grid
    _put(ws, "AJ3", "#Lookup")
    for col, h in zip(("AJ", "AK", "AL"), ("K1", "K2", "Out")):
        _put(ws, f"{col}4", h)
    for i, row in enumerate((("A", "B", 1), ("A", "C", 2)), start=5):
        for col, v in zip(("AJ", "AK", "AL"), row):
            _put(ws, f"{col}{i}", v)
    _put(ws, "AN3", "#Read")
    _put(ws, "AN4", "v")
    _put(ws, "AN5", '=MM_READTABLENAN(AJ4:AL6,"Out","A","C")')
    _put(ws, "AP3", "#Inverse")
    _put(ws, "AP4", "=MM_RANGE(MINVERSE(MM_TABLE(AL5)))")

    settings = wb.create_sheet("Settings")
    _put(settings, "A1", "#Settings /ProjectSettings")
    for col, h in zip("ABCD", ("Name", "Value", "Locked", "Hidden")):
        _put(settings, f"{col}2", h)
    for i, row in enumerate((("PerformanceProfiler", False, False, True), ("StopOnFirstNaN", False, False, False)), start=3):
        for col, v in zip("ABCD", row):
            _put(settings, f"{col}{i}", v)

    out = wb.create_sheet("Outputs")
    _put(out, "A3", "#Results /Export")
    _put(out, "A4", "Year")
    _put(out, "B4", "Amount")
    for i, (y, a) in enumerate(((2024, 100), (2025, 200)), start=5):
        _put(out, f"A{i}", y)
        _put(out, f"B{i}", a)
    _put(out, "D3", "#ExpSet /ExportSettings")
    for col, h in zip("DEFG", ("GridName", "ExportByInstance", "FileNamePattern", "FileExtension")):
        _put(out, f"{col}4", h)
    for col, v in zip("DEFG", ("Results", "true", "{GridName}_{yyyyMMdd}", ".csv")):
        _put(out, f"{col}5", v)

    path = tmp_path / "flagged_model.xlsx"
    wb.save(path)
    wb.close()
    return path


@pytest.fixture
def flagged_model_broken_xlsx(tmp_path: Path) -> Path:
    """The same features, each broken the way its KB article says not to."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    _put(ws, "A1", LABEL)

    # /Reorder without /Input, duplicate headers, unknown flag, mixed header row
    _put(ws, "A3", "#Assumptions /Reorder /Inpt")
    for col, h in zip("ABC", ("Key", "Key", 42)):
        _put(ws, f"{col}4", h)
    for i in range(5, 8):
        _put(ws, f"A{i}", i)
        _put(ws, f"B{i}", i)
        _put(ws, f"C{i}", i)

    # /Parameters with a wrong header, a bad type and a bad range
    _put(ws, "E3", "#Params /Parameters")
    for col, h in zip("EFGH", ("Label", "Kind", "PossibleValues", "Values")):
        _put(ws, f"{col}4", h)
    for i, row in enumerate((("Rate", "numbre", "1|0", 0.5), ("Mode", "dropdown", "", "A")), start=5):
        for col, v in zip("EFGH", row):
            _put(ws, f"{col}{i}", v)

    # MM_RESULT on a loop with the wrong case + the same loop twice
    _put(ws, "J3", "#Calc")
    _put(ws, "J4", "Loop")
    _put(ws, "J5", '=MM_LOOP("Scenario", A5:A7)')
    _put(ws, "J6", '=MM_RESULT(J5,"scenario",1)')
    _put(ws, "J7", '=MM_RESULT(J5,"Scenario",1,"Scenario",2)')

    # MM_SETSIZE nested in a function and with an occupied destination
    _put(ws, "L3", "#Table /Resize.solo")
    _put(ws, "L4", "Val")
    _put(ws, "L5", "=SUM(L4,MM_SETSIZE(3,1))")
    _put(ws, "L6", "leftover")
    _put(ws, "L7", "leftover")

    # bad /Group, backup without destination + button inside the source grid, HideRows with text
    _put(ws, "N3", "#G1 /Group.(Block).x.0")
    _put(ws, "N4", "h")
    _put(ws, "N5", 1)
    _put(ws, "P3", "#Src /BackupSource.(b1)")
    _put(ws, "P4", "h")
    _put(ws, "P5", '=MM_BACKUPBUTTON("inside a backup grid")')
    _put(ws, "R3", "#Hide")
    _put(ws, "R4", "h")
    _put(ws, "S4", "/HideRows")
    _put(ws, "R5", 1)
    _put(ws, "S5", "maybe")

    # NbSimulations as text, Translations with a bad code, MM_RANGE not alone
    _put(ws, "U3", "#NbSimulations")
    _put(ws, "U4", "lots")
    _put(ws, "W3", "#Tr /Translations")
    _put(ws, "W4", "en")
    _put(ws, "X4", "english")
    _put(ws, "W5", "Hello")
    _put(ws, "X5", "Hello")
    _put(ws, "Z3", "#Inverse")
    _put(ws, "Z4", "v")
    _put(ws, "Z5", "=MM_RANGE(MINVERSE(MM_TABLE(A5)))")
    _put(ws, "AA5", 1)

    # MM_READTABLE whose Header matches no column and whose range skips the header row
    _put(ws, "AC3", "#Lookup")
    for col, h in zip(("AC", "AD", "AE"), ("K1", "K2", "Out")):
        _put(ws, f"{col}4", h)
    for i, row in enumerate((("A", "B", 1), ("A", "B", 2)), start=5):
        for col, v in zip(("AC", "AD", "AE"), row):
            _put(ws, f"{col}{i}", v)
    _put(ws, "AG3", "#Read")
    _put(ws, "AG4", "v")
    _put(ws, "AG5", '=MM_READTABLE(AC5:AE6,"Result","A","B")')

    # a '#Title' trapped inside a grid (no empty row between two grids)
    _put(ws, "AI3", "#First")
    _put(ws, "AI4", "h")
    _put(ws, "AI5", 1)
    _put(ws, "AI6", "#Second")
    _put(ws, "AI7", "h")
    _put(ws, "AI8", 2)

    # two /ProjectSettings grids with the wrong columns / values
    settings = wb.create_sheet("Settings")
    _put(settings, "A1", "#Settings /ProjectSettings")
    for col, h in zip("ABCD", ("Setting", "Value", "Locked", "Hidden")):
        _put(settings, f"{col}2", h)
    _put(settings, "A3", "X")
    _put(settings, "B3", True)
    _put(settings, "F1", "#Settings2 /ProjectSettings")
    for col, h in zip("FGHI", ("Name", "Value", "Locked", "Hidden")):
        _put(settings, f"{col}2", h)
    _put(settings, "F3", "Y")
    _put(settings, "H3", "sometimes")
    _put(settings, "A4", "EnableDebugMode")  # PRJ-006: a setting name Mind's converter rejects

    # /ExportSettings with an unknown column, a non-export GridName and a bad boolean
    out = wb.create_sheet("Outputs")
    _put(out, "A3", "#Results /Export")
    _put(out, "A4", "Year")
    _put(out, "A5", 2024)
    _put(out, "D3", "#ExpSet /ExportSettings")
    for col, h in zip("DEFG", ("GridName", "ExportByInstance", "FileNamePattern", "Colour")):
        _put(out, f"{col}4", h)
    for col, v in zip("DEFG", ("Nope", "perhaps", "{GridName}_{Bogus}", "red")):
        _put(out, f"{col}5", v)

    path = tmp_path / "flagged_model_broken.xlsx"
    wb.save(path)
    wb.close()
    return path


@pytest.fixture
def array_formula_xlsx(tmp_path: Path) -> Path:
    """An array formula whose target range spills past its grid, plus an
    '@' implicit-intersection formula, an unsupported native function and a
    name that is no Excel function at all."""
    from openpyxl.worksheet.formula import ArrayFormula

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Arr"
    _put(ws, "A1", LABEL)
    _put(ws, "A3", "#Grid")
    _put(ws, "A4", "h")
    _put(ws, "A5", 1)
    _put(ws, "A6", 2)
    _put(ws, "C3", "#Spill")
    _put(ws, "C4", "h")
    ws["C5"] = ArrayFormula("C5:C9", "=A5:A6*2")  # grid C4:C5 but the array targets C5:C9
    _put(ws, "E3", "#Other")
    _put(ws, "E4", "h")
    _put(ws, "E5", "=@INDEX(A5:A6,1)")
    _put(ws, "E6", "=_xlfn.XLOOKUP(1,A5:A6,A5:A6)")
    _put(ws, "E7", "=MyMacroFunction(A5)")
    path = tmp_path / "array_formula.xlsx"
    wb.save(path)
    wb.close()
    return path


@pytest.fixture
def untitled_grids_xlsx(tmp_path: Path) -> Path:
    """Untitled grids with the kinds of context the title action reads:
    a label above with a blank row between, a caption directly above a
    two-column table (which splits it into two blocks), a label to the left,
    a table with only a text header row, and a grid starting on row 1."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    _put(ws, "A1", LABEL)  # standalone, far from any grid -> reported, left alone
    _put(ws, "B2", "Premium table")  # label above (blank row 3 between)
    for col, h in zip("BCD", ("Key", "Value", "Note")):
        _put(ws, f"{col}4", h)
    for i, (k, v) in enumerate((("alpha", 1), ("beta", 2)), start=5):
        _put(ws, f"B{i}", k)
        _put(ws, f"C{i}", v)
        _put(ws, f"D{i}", "n")
    _put(ws, "F4", "Rates")  # caption directly above F5:G7
    for i, (a, b) in enumerate(((1, 2), (3, 4), (5, 6)), start=5):
        _put(ws, f"F{i}", a)
        _put(ws, f"G{i}", b)
    _put(ws, "J4", "Factors")  # label to the left of L4:M5 (K empty)
    for i, (a, b) in enumerate(((7, 8), (9, 10)), start=4):
        _put(ws, f"L{i}", a)
        _put(ws, f"M{i}", b)
    _put(ws, "O4", "Year")  # header-row-only context
    _put(ws, "P4", "Amount")
    _put(ws, "O5", 2024)
    _put(ws, "P5", 100)
    top = wb.create_sheet("Top")
    _put(top, "A1", "X")  # grid starting on row 1
    _put(top, "B1", "Y")
    _put(top, "A2", 1)
    _put(top, "B2", 2)
    path = tmp_path / "untitled_grids.xlsx"
    wb.save(path)
    wb.close()
    return path


@pytest.fixture
def section_heading_grids_xlsx(tmp_path: Path) -> Path:
    """The layout real models use for a run of calculation blocks: a section
    number beside a section title, directly above the block it heads. The cell
    above each block is therefore *occupied* (it is another grid), which is the
    case the title action used to skip. Two blocks, so the second one also
    exercises the post-insert coordinate shift."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CF"
    _put(ws, "B2", 1)
    _put(ws, "C2", "Demographic Assumptions")
    for i, (k, v) in enumerate((("Gender", "M"), ("Smoker", "NS"), ("Age", 45)), start=3):
        _put(ws, f"C{i}", k)
        _put(ws, f"D{i}", v)
    _put(ws, "B8", 2)
    _put(ws, "C8", "Economic Assumptions")
    for i, (k, v) in enumerate((("CPI Base", 100), ("CPI Current", 103)), start=9):
        _put(ws, f"C{i}", k)
        _put(ws, f"D{i}", v)
    path = tmp_path / "section_headings.xlsx"
    wb.save(path)
    wb.close()
    return path
