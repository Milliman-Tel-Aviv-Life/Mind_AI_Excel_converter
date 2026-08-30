"""1.6.8: regression tests for the defects the adversarial review of
app/mind_loop.py found before its first unattended run -- each one a way the
loop could have converged on false evidence, crashed, or (delete) acted on
the wrong Mind project."""
from __future__ import annotations

from pathlib import Path

import openpyxl

from app import mind_loop as ml


def _book(path: Path, cells: dict[str, dict[str, object]]) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, values in cells.items():
        ws = wb.create_sheet(sheet)
        for ref, v in values.items():
            ws[ref] = v
    wb.save(path)
    wb.close()
    return path


def _plan():
    return [
        {"id": "create_grid_titles", "default_on": True, "count": 3, "rule_ids": ["STR-004"], "operations": [{"op": "insert_row", "sheet": "CF", "row": 3}, {"op": "set_value", "sheet": "CF", "cell": "C3"}]},
        {"id": "separate_merged_grids", "default_on": True, "count": 1, "rule_ids": ["STR-001"], "operations": [{"op": "insert_row", "sheet": "CF", "row": 10}, {"op": "insert_row", "sheet": "CF", "row": 20}]},
        {"id": "fix_broken_refs", "default_on": False, "count": 2, "rule_ids": ["REF-001"], "operations": [{"op": "set_formula", "sheet": "CF", "cell": "J5"}]},
    ]


def test_rewritten_formula_cells_are_compared_under_the_error_only_rule(tmp_path):
    """fix_broken_refs rewrites #REF! -> NA(): the rewritten cell may keep its
    value or turn an error into another error, but a valid value that moves is
    a real difference (IF(ISNA(BrokenName),...) flips 'present' -> 'missing')."""
    source = _book(tmp_path / "s.xlsx", {"S": {"A1": "present", "A2": "#REF!", "A3": 4, "A4": True, "B1": 7}})
    prepared = _book(tmp_path / "p.xlsx", {"S": {"A1": "missing", "A2": "#N/A", "A3": 7, "A4": False, "B1": 7}})
    applied = [{"op": "set_formula", "action_id": "fix_broken_refs", "rule_id": "REF-001", "sheet": "S", "cell": c} for c in ("A1", "A2", "A3", "A4")]
    out = ml.compare_values(source, prepared, applied, formulas=[], defined_names=[])
    assert sorted(d["cell"] for d in out["listed"]) == ["A1", "A3", "A4"]
    assert all(d["rewritten"] for d in out["listed"]) and not out["match"]
    # a content write (a title) on the same cell is the plan's intent and is skipped
    applied2 = applied + [{"op": "set_value", "action_id": "create_grid_titles", "rule_id": "STR-004", "sheet": "S", "cell": "A1", "after": "#T"}]
    assert "A1" not in {d["cell"] for d in ml.compare_values(source, prepared, applied2, [], [])["listed"]}


def test_compare_values_allows_titles_written_after_inserts_on_a_renamed_sheet(tmp_path):
    # hide_marker renames the sheet last; the title op is keyed by the source name
    source = _book(tmp_path / "s.xlsx", {"Hid": {"A1": "h", "A2": 1}})
    prepared = _book(tmp_path / "p.xlsx", {"Hid &&Hide": {"A1": "#T", "A2": "h", "A3": 1}})
    applied = [
        {"op": "rename_sheet", "action_id": "hide_marker", "rule_id": "STR-007", "sheet": "Hid", "before": "Hid", "after": "Hid &&Hide"},
        {"op": "insert_row", "action_id": "create_grid_titles", "rule_id": "STR-004", "sheet": "Hid", "row": 1},
        {"op": "set_value", "action_id": "create_grid_titles", "rule_id": "STR-004", "sheet": "Hid", "cell": "A1", "after": "#T", "after_inserts": True, "insert_at": 1},
    ]
    out = ml.compare_values(source, prepared, applied, formulas=[], defined_names=[])
    assert out["match"] and out["differences"] == 0 and out["inserted_cells_checked"] == 1


def test_compare_values_ignores_chart_sheets(tmp_path):
    from openpyxl.chart import BarChart, Reference

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=1, min_row=1, max_row=1))
    wb.create_chartsheet("Chart").add_chart(chart)  # a real chart sheet carries a drawing part
    wb.save(tmp_path / "s.xlsx")
    wb.save(tmp_path / "p.xlsx")
    wb.close()
    out = ml.compare_values(tmp_path / "s.xlsx", tmp_path / "p.xlsx", [], [], [])
    assert out["match"] and out["compared"] == 1


def test_blame_follows_precedents_to_the_formula_rewrite_not_the_row_insert():
    plan = [
        {"id": "fix_broken_refs", "default_on": False, "count": 1, "rule_ids": ["REF-001"], "operations": [{"op": "set_formula", "sheet": "CF", "cell": "J5"}]},
        {"id": "create_grid_titles", "default_on": True, "count": 1, "rule_ids": ["STR-004"], "operations": [{"op": "insert_row", "sheet": "CF", "row": 30}]},
    ]
    formulas = [
        {"sheet": "CF", "cell": "J6", "formula": "=J5*2"},
        {"sheet": "Out", "cell": "B2", "formula": "=Total+1"},
        {"sheet": "CF", "cell": "K9", "formula": "=SUM(J1:J8)"},
    ]
    names = [{"name": "Total", "value": "CF!$K$9"}]
    assert ml.actions_touching(plan, [{"sheet": "CF", "cell": "J6"}], formulas, names) == ["fix_broken_refs"]
    assert ml.actions_touching(plan, [{"sheet": "Out", "cell": "B2"}], formulas, names) == ["fix_broken_refs"]  # via a name, then a range
    # nothing reachable -> the row insert on that sheet takes the blame
    assert ml.actions_touching(plan, [{"sheet": "CF", "cell": "Z99"}], formulas, names) == ["create_grid_titles"]
    it = {"apply": {"status": "APPLIED"}, "names": {"fragmented": False}, "numbers": {"ran": True, "match": False, "differences": 1, "listed": [{"sheet": "CF", "cell": "J6", "prepared_cell": "J6"}]}}
    d = ml.decide(it, plan, enabled={"fix_broken_refs"}, disabled=set(), formulas=formulas, defined_names=names)
    assert d["verdict"] == "retry" and d["disable"] == ["fix_broken_refs"]


def test_a_gate_that_could_not_run_is_never_a_pass():
    good_mind = {"skipped": False, "convert": {"success": True}, "run": {"completed": True, "audit_consistent": True}, "counts": {"mind_untitled": 0, "app_untitled": 0}}
    it = {"apply": {"status": "APPLIED"}, "names": {"fragmented": False}, "mind": good_mind, "numbers": {"ran": False, "requested": True, "message": "Excel COM failed"}}
    d = ml.decide(it, _plan(), set(), set())
    assert d["verdict"] == "stuck" and "could not be verified" in d["reason"]
    off = {**it, "numbers": {"ran": False, "requested": False, "message": "disabled"}}
    assert ml.decide(off, _plan(), set(), set())["verdict"] == "converged"


def test_mind_failures_after_convert_are_stuck_not_converged():
    base = {"apply": {"status": "APPLIED"}, "numbers": {"ran": True, "match": True}, "names": {"fragmented": False}}
    add_failed = {**base, "mind": {"skipped": False, "convert": {"success": True, "errors": []}, "error": "Add template to Milliman Mind did not lead to a Run button"}}
    d = ml.decide(add_failed, _plan(), set(), set())
    assert d["verdict"] == "stuck" and "Add template" in d["reason"]
    not_logged_in = {**base, "mind": {"skipped": False, "error": "could not reach the Project manager (not logged in?)"}}
    assert "Project manager" in ml.decide(not_logged_in, _plan(), set(), set())["reason"]
    never_ran = {**base, "run_model": True, "mind": {"skipped": False, "convert": {"success": True}, "counts": {"mind_untitled": 0, "app_untitled": 0}}}
    assert ml.decide(never_ran, _plan(), set(), set())["verdict"] == "stuck"
    convert_only = {**never_ran, "run_model": False}
    assert ml.decide(convert_only, _plan(), set(), set())["verdict"] == "converged"


def test_run_loop_reports_a_verdict_when_prep_cannot_be_applied(tmp_path, section_heading_grids_xlsx, monkeypatch):
    def partial(source, work_dir, ops, prefer_excel=True):
        work_dir.mkdir(parents=True, exist_ok=True)
        return {"status": "PARTIAL", "output_path": work_dir / "x.xlsx", "applied": [], "failed": [ops[0]], "message": "1 of 3 operations failed: part of an array", "verified_opens_in_excel": True}

    monkeypatch.setattr(ml.prep, "apply_operations", partial)
    events: list[dict] = []
    cfg = ml.LoopConfig(source=section_heading_grids_xlsx, work_dir=tmp_path / "loop", max_iterations=2, use_assistant=False, skip_mind=True)
    report = ml.run_loop(cfg, events.append)
    assert report["verdict"] == "stuck" and "prep could not be applied" in report["reason"]
    assert any(e["event"] == "decision" for e in events) and events[-1]["event"] == "done"
    assert report["iterations"][0]["finished"]
