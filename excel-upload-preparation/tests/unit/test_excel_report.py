from pathlib import Path

import openpyxl
import pytest

from app.config import load_config
from app.excel_com import com_available, verify_opens_in_excel
from app.excel_report import HEADERS, REPORT_SHEET_NAME, SUMMARY_SHEET_NAME, build_report_workbook, build_standalone_report
from app.modes import plan_mode
from app.rules_engine import RulesEngine

needs_excel = pytest.mark.skipif(not com_available(), reason="requires pywin32 + an installed Excel")


def _run(tmp_path, source):
    config = load_config()
    engine = RulesEngine()
    result = plan_mode.run(source, tmp_path / "work", config, engine=engine)
    return result, Path(result["workbook_analysis"]["source"]["copy_path"]), engine


def test_standalone_report_is_a_valid_workbook_with_all_findings(tmp_path, loops_broken_xlsx):
    result, _, engine = _run(tmp_path, loops_broken_xlsx)
    output = build_standalone_report(result["validation_report"], tmp_path / "standalone.xlsx", rules_by_id=engine.rules, source_name="x.xlsx")
    wb = openpyxl.load_workbook(output)
    assert wb.sheetnames == ["Summary", "Findings"]
    ws = wb["Findings"]
    assert [c.value for c in ws[1]] == HEADERS
    assert ws.max_row == result["validation_report"]["summary"]["finding_count"] + 1
    assert wb["Summary"]["B4"].value == result["validation_report"]["status"]


def test_openpyxl_fallback_appends_sheets_and_says_so(tmp_path, plain_grid_xlsx):
    result, copy_path, engine = _run(tmp_path, plain_grid_xlsx)
    built = build_report_workbook(copy_path, result["validation_report"], tmp_path / "report.xlsx", rules_by_id=engine.rules, prefer_excel=False, verify=False)
    assert built.method == "openpyxl"
    assert any("openpyxl" in w for w in built.warnings)
    wb = openpyxl.load_workbook(built.path)
    assert REPORT_SHEET_NAME in wb.sheetnames and SUMMARY_SHEET_NAME in wb.sheetnames
    assert wb["Data"]["A2"].value == "Header1"  # the original sheet, untouched
    assert [c.value for c in wb[REPORT_SHEET_NAME][1]] == HEADERS
    # the readiness sheets are hidden so an upload of this copy ignores them; a visible sheet stays active
    assert wb[REPORT_SHEET_NAME].sheet_state == "hidden" and wb[SUMMARY_SHEET_NAME].sheet_state == "hidden"
    assert wb.active.sheet_state == "visible"


@needs_excel
def test_excel_written_report_opens_in_excel(tmp_path, loops_broken_xlsx):
    """The 1.3.0 guarantee: the downloadable workbook copy is written by Excel
    and re-opened by Excel to prove it is a real, working file."""
    result, copy_path, engine = _run(tmp_path, loops_broken_xlsx)
    built = build_report_workbook(copy_path, result["validation_report"], tmp_path / "report.xlsx", rules_by_id=engine.rules)
    assert built.method == "excel_com"
    assert built.verified_opens_in_excel is True
    assert built.warnings == []
    wb = openpyxl.load_workbook(built.path)
    assert wb.sheetnames[0] == SUMMARY_SHEET_NAME and wb.sheetnames[-1] == REPORT_SHEET_NAME
    ws = wb[REPORT_SHEET_NAME]
    assert [c.value for c in ws[1]] == HEADERS
    assert ws.max_row == result["validation_report"]["summary"]["finding_count"] + 1
    assert ws.sheet_state == "hidden" and wb[SUMMARY_SHEET_NAME].sheet_state == "hidden" and wb.active.sheet_state == "visible"
    assert verify_opens_in_excel(built.path)["opens"] is True
