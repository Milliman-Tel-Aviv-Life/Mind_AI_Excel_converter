"""1.6.8: the autonomous run-in-Mind loop (app/mind_loop.py) -- every decision
it makes is a pure function tested here without Excel, Mind or the assistant.
The browser operations in app/mind_client.py are exercised only by the live
loop; this file covers what turns their results into the next iteration."""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
import pytest

from app import mind_loop as ml
from app.mind_client import untitled_entries


# --- coordinates and value equivalence ------------------------------------------------


def test_row_map_pushes_rows_at_or_below_each_insert():
    f = ml.row_map([3, 9])
    assert [f(r) for r in (1, 2, 3, 4, 8, 9, 10)] == [1, 2, 4, 5, 9, 11, 12]
    assert ml.row_map([])(7) == 7


def test_same_value_rules():
    assert ml._same(1.0, 1.0 + 1e-12)
    assert not ml._same(1.0, 1.001)
    assert ml._same("#REF!", "#N/A")  # an error may become another error
    assert not ml._same("#REF!", 5)   # but never a value
    assert not ml._same(None, 0)
    assert ml._same(dt.datetime(2026, 1, 1), dt.datetime(2026, 1, 1))
    assert not ml._same(True, 1) or ml._same(True, True)


def test_volatile_cells_follow_references_and_defined_names():
    formulas = [
        {"sheet": "Run", "cell": "P3", "formula": "=NOW()"},
        {"sheet": "Run", "cell": "Q3", "formula": "=P3+1"},
        {"sheet": "Calc", "cell": "A1", "formula": "=Run!Q3*2"},
        {"sheet": "Calc", "cell": "A2", "formula": "=TimeNow+0"},
        {"sheet": "Calc", "cell": "B1", "formula": "=SUM(Data!A:A)"},  # whole column, not volatile
        {"sheet": "Calc", "cell": "C1", "formula": "=1+1"},
    ]
    names = [{"name": "TimeNow", "value": "Run!$P$3"}]
    vol = ml.volatile_cells(formulas, names)
    assert vol == {("Run", "P3"), ("Run", "Q3"), ("Calc", "A1"), ("Calc", "A2")}


def _book(path: Path, cells: dict[str, dict[str, object]]) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, values in cells.items():
        ws = wb.create_sheet(sheet)
        for ref, v in values.items():
            ws[ref] = v
    wb.save(path)
    wb.close()
    return path


def test_compare_values_maps_inserts_allows_plan_cells_and_reports_real_changes(tmp_path):
    source = _book(tmp_path / "s.xlsx", {"CF": {"B2": 1, "C2": "Section", "C3": "Gender", "D3": "M", "C4": "Age", "D4": 45, "F9": 100.0, "H1": "#REF!"}})
    # the plan inserted a row at 3 and wrote '#Demographics' at C3 (post-insert); everything else shifted down one
    prepared = _book(tmp_path / "p.xlsx", {"CF": {"B2": 1, "C2": "Section", "C3": "#Demographics", "C4": "Gender", "D4": "M", "C5": "Age", "D5": 45, "F10": 100.5, "H1": "#N/A"}})
    applied = [
        {"op": "insert_row", "action_id": "create_grid_titles", "rule_id": "STR-004", "sheet": "CF", "row": 3},
        {"op": "set_value", "action_id": "create_grid_titles", "rule_id": "STR-004", "sheet": "CF", "cell": "C3", "after": "#Demographics", "after_inserts": True, "insert_at": 3},
    ]
    out = ml.compare_values(source, prepared, applied, formulas=[], defined_names=[])
    assert out["compared"] >= 6
    # the only genuine change: F9 (-> F10) went from 100.0 to 100.5; the error->error in H1 is allowed
    assert out["differences"] == 1 and out["listed"][0]["cell"] == "F9" and out["listed"][0]["prepared_cell"] == "F10"
    assert not out["match"]


def test_compare_values_flags_unexpected_content_in_an_inserted_row(tmp_path):
    source = _book(tmp_path / "s.xlsx", {"S": {"A1": "x", "A2": "y"}})
    prepared = _book(tmp_path / "p.xlsx", {"S": {"A1": "x", "A2": "#Title", "B2": "stray", "A3": "y"}})
    applied = [
        {"op": "insert_row", "action_id": "create_grid_titles", "rule_id": "STR-004", "sheet": "S", "row": 2},
        {"op": "set_value", "action_id": "create_grid_titles", "rule_id": "STR-004", "sheet": "S", "cell": "A2", "after": "#Title", "after_inserts": True, "insert_at": 2},
    ]
    out = ml.compare_values(source, prepared, applied, formulas=[], defined_names=[])
    assert out["differences"] == 1 and out["listed"][0]["prepared_cell"] == "B2"


def test_compare_values_skips_volatile_cells(tmp_path):
    source = _book(tmp_path / "s.xlsx", {"Run": {"P3": dt.datetime(2026, 8, 30, 1, 0), "Q3": 5}})
    prepared = _book(tmp_path / "p.xlsx", {"Run": {"P3": dt.datetime(2026, 8, 30, 2, 0), "Q3": 5}})
    formulas = [{"sheet": "Run", "cell": "P3", "formula": "=NOW()"}]
    out = ml.compare_values(source, prepared, [], formulas, [])
    assert out["match"] and out["volatile_skipped"] == 1


# --- decisions -----------------------------------------------------------------------


def _plan(**counts):
    return [
        {"id": "create_grid_titles", "default_on": True, "count": counts.get("titles", 3), "rule_ids": ["STR-004"], "operations": [{"op": "insert_row", "sheet": "CF", "row": 3}, {"op": "set_value", "sheet": "CF", "cell": "C3"}]},
        {"id": "separate_merged_grids", "default_on": True, "count": counts.get("separate", 1), "rule_ids": ["STR-001"], "operations": [{"op": "insert_row", "sheet": "CF", "row": 10}, {"op": "insert_row", "sheet": "CF", "row": 20}]},
        {"id": "fix_broken_refs", "default_on": False, "count": counts.get("refs", 2), "rule_ids": ["REF-001"], "operations": [{"op": "set_formula", "sheet": "CF", "cell": "J5"}]},
        {"id": "explicit_colors", "default_on": True, "count": 0, "rule_ids": ["FMT-002"], "operations": []},
    ]


def test_auto_enable_turns_on_opt_in_actions_whose_rule_is_an_error():
    report = {"findings": [{"rule_id": "REF-001", "status": "ERROR"}, {"rule_id": "FMT-005", "status": "WARNING"}]}
    assert ml.auto_enable(_plan(), report) == ["fix_broken_refs"]
    assert ml.auto_enable(_plan(), {"findings": [{"rule_id": "REF-001", "status": "WARNING"}]}) == []


def test_decide_enables_the_fix_mind_asks_for_then_gives_up_when_it_has_none():
    it = {"apply": {"status": "APPLIED"}, "numbers": {"ran": True, "match": True}, "names": {"fragmented": False},
          "mind": {"skipped": False, "convert": {"success": False, "errors": ["Formula compilation error: #REF! in Cashflows!J5"]}}}
    d = ml.decide(it, _plan(), enabled=set(), disabled=set())
    assert d["verdict"] == "retry" and d["enable"] == ["fix_broken_refs"]
    d2 = ml.decide(it, _plan(), enabled={"fix_broken_refs"}, disabled=set())
    assert d2["verdict"] == "stuck" and "no fix" in d2["reason"]
    it3 = {**it, "mind": {"skipped": False, "convert": {"success": False, "errors": ["The EnableDebugMode setting does not exist."]}}}
    assert ml.decide(it3, _plan(), set(), set())["verdict"] == "stuck"


def test_decide_disables_the_action_that_moved_a_number():
    it = {"apply": {"status": "APPLIED"}, "names": {"fragmented": False},
          "numbers": {"ran": True, "match": False, "differences": 1, "listed": [{"sheet": "CF", "cell": "J5", "prepared_cell": "J5"}]}}
    d = ml.decide(it, _plan(), enabled={"fix_broken_refs"}, disabled=set())
    assert d["verdict"] == "retry" and d["disable"] == ["fix_broken_refs"]
    # a diff nobody wrote to -> stuck, never a blind retry
    it2 = {**it, "numbers": {"ran": True, "match": False, "differences": 1, "listed": [{"sheet": "Other", "cell": "Z9", "prepared_cell": "Z9"}]}}
    assert ml.decide(it2, _plan(), {"fix_broken_refs"}, set())["verdict"] == "stuck"


def test_decide_disables_the_insert_heaviest_action_on_fragmentation():
    it = {"apply": {"status": "APPLIED"}, "numbers": {"ran": True, "match": True},
          "names": {"fragmented": True, "grids": 648, "baseline_grids": 156}}
    # default-on actions are active without being in `enabled` -- and can be disabled
    d = ml.decide(it, _plan(), enabled=set(), disabled=set())
    assert d["verdict"] == "retry" and d["disable"] == ["separate_merged_grids"]
    d2 = ml.decide(it, _plan(), enabled=set(), disabled={"separate_merged_grids"})
    assert d2["verdict"] == "retry" and d2["disable"] == ["create_grid_titles"]


def test_decide_converges_only_when_mind_agrees():
    base = {"apply": {"status": "APPLIED"}, "numbers": {"ran": True, "match": True}, "names": {"fragmented": False}}
    good = {**base, "mind": {"skipped": False, "convert": {"success": True}, "run": {"completed": True, "audit_consistent": True}, "counts": {"mind_untitled": 36, "app_untitled": 36}}}
    assert ml.decide(good, _plan(), set(), set())["verdict"] == "converged"
    # Mind reading a block differently from the app is a detection gap: reported, never a block
    off = {**good, "mind": {**good["mind"], "counts": ml.untitled_counts(["D13", "D25", "H60"], ["C13", "C25", "H60"])}}
    d = ml.decide(off, _plan(), set(), set())
    assert d["verdict"] == "converged" and "2 Mind-only, 2 app-only" in d["reason"]
    c = ml.untitled_counts(["D13", "H60", "H60"], ["H60", "C13"])
    assert c == {"mind_untitled": 3, "app_untitled": 2, "unexplained": ["D13", "H60"], "app_only": ["C13"]}
    no_audit = {**good, "mind": {**good["mind"], "run": {"completed": True, "audit_consistent": False}}}
    assert ml.decide(no_audit, _plan(), set(), set())["verdict"] == "stuck"
    local_only = {**base, "mind": {"skipped": True, "reason": "skip_mind"}}
    assert ml.decide(local_only, _plan(), set(), set())["verdict"] == "converged"


def test_untitled_entries_decode_mind_labels_to_cells():
    out = untitled_entries(["work", "Information", "Untitled(60,8)", "Untitled(1210,80)", "Not Untitled"])
    assert [(u["label"], u["cell"]) for u in out] == [("Untitled(60,8)", "H60"), ("Untitled(1210,80)", "CB1210")]


# --- the loop itself, offline --------------------------------------------------------


@pytest.mark.skipif(not __import__("app.excel_com", fromlist=["com_available"]).com_available(), reason="needs Excel")
def test_run_loop_offline_converges_on_local_gates(tmp_path, section_heading_grids_xlsx):
    """skip_mind + no assistant: prep through Excel, recalc both copies, compare,
    and converge on the local gates alone. Proves the orchestration end to end
    without touching Mind."""
    events: list[dict] = []
    cfg = ml.LoopConfig(source=section_heading_grids_xlsx, work_dir=tmp_path / "loop", max_iterations=2, use_assistant=False, skip_mind=True)
    report = ml.run_loop(cfg, events.append)
    assert report["verdict"] == "converged", report
    it = report["iterations"][0]
    assert it["apply"]["status"] == "APPLIED"
    assert it["numbers"]["ran"] and it["numbers"]["match"], it["numbers"]
    assert it["names"]["still_titleable"] == 0 and not it["names"]["fragmented"]
    assert (tmp_path / "loop" / "loop_report.json").is_file()
    assert any(e["event"] == "done" for e in events)
    assert Path(report["final_workbook"]).is_file()
