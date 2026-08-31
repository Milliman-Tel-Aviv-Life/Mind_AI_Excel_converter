"""1.7.0 Grid Namer: the user selects an area of the sheet, chooses a name and
flags, and the app writes the '#Name /Flags' title for the grid Mind's
detection sees there -- with the same reference-safety rules as the automatic
titler, which then names everything the user did not touch without colliding
with the user's choices."""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import prep
from app.inventory import build_analysis
from app.web import server


def _wb(path: Path, sheets: dict[str, dict[str, object]]) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, cells in sheets.items():
        ws = wb.create_sheet(name)
        for ref, v in cells.items():
            ws[ref] = v
    wb.save(path)
    wb.close()
    return path


def _analysis(tmp_path: Path, sheets: dict[str, dict[str, object]]) -> dict:
    return build_analysis(_wb(tmp_path / "m.xlsx", sheets), tmp_path / "w", "t")


def test_named_area_titles_an_untitled_grid_with_canonical_flags(tmp_path):
    """A selection anywhere inside the grid resolves to it; flags get the
    documented casing; the title lands in the free cell above."""
    analysis = _analysis(tmp_path, {"Data": {"B3": "Age", "C3": "Rate", "B4": 30, "C4": 0.01, "B5": 31, "C5": 0.02}})
    out = prep.plan_named_areas(analysis, [{"sheet": "Data", "ref": "B4:C4", "name": "Mortality Rates", "flags": ["input", "/Export"]}])
    assert [o["op"] for o in out["ops"]] == ["set_value"]
    op = out["ops"][0]
    assert op["cell"] == "B2" and op["after"] == "#Mortality Rates /Input /Export"
    assert out["named"] == ["Data!B3:C5"] and out["reserved"] == ["Mortality Rates"]
    assert "Data!B3:C5" in out["exclude"] and ("Data", "B2") in out["titled_cells"]


def test_selection_across_two_separate_grids_is_refused(tmp_path):
    analysis = _analysis(tmp_path, {
        "S": {"B3": "H", "C3": "K", "B4": 1, "C4": 2, "F3": "H2", "G3": "K2", "F4": 1, "G4": 2},
    })
    out = prep.plan_named_areas(analysis, [{"sheet": "S", "ref": "B3:G4", "name": "Everything", "flags": []}])
    assert not out["ops"]
    assert any("touches 2 grids" in s and "select and name one at a time" in s for s in out["skipped"]), out["skipped"]


def test_caption_pair_is_named_as_one_grid(tmp_path):
    """The documented caption-above-a-table split (two detected grids) is the
    one two-grid selection that works: the caption becomes the title."""
    analysis = _analysis(tmp_path, {"S": {"B2": "Rates", "B3": 1, "C3": 2, "B4": 3, "C4": 4}})
    out = prep.plan_named_areas(analysis, [{"sheet": "S", "ref": "B2:C4", "name": "Lapse Rates", "flags": ["Input"]}])
    assert [o["op"] for o in out["ops"]] == ["set_value"]
    op = out["ops"][0]
    assert op["cell"] == "B2" and op["after"] == "#Lapse Rates /Input" and op["before"] == "Rates"
    assert len(out["named"]) == 1 and len(out["exclude"]) == 2  # the table grid is excluded from the automatic pass too


def test_selecting_only_the_table_of_a_caption_pair_names_the_pair(tmp_path):
    """Found live on the caption layout: naming only the table half leaves the
    caption column to the automatic pass, whose separate title lands adjacent
    and the re-detected grids merge wrongly. One half of an untitled caption
    pair is therefore extended to the pair -- the caption becomes the title --
    and the automatic pass writes nothing on that sheet."""
    analysis = _analysis(tmp_path, {
        "S": {"A1": "Inputs", "A2": "Order", "A3": "r1", "B2": "Item1", "C2": "Item2", "B3": 1, "C3": 2},
    })
    out = prep.plan_named_areas(analysis, [{"sheet": "S", "ref": "B2:C3", "name": "Projection Table", "flags": ["Input"]}])
    assert [o["op"] for o in out["ops"]] == ["set_value"]
    assert out["ops"][0]["cell"] == "A1" and out["ops"][0]["after"] == "#Projection Table /Input" and out["ops"][0]["before"] == "Inputs"
    assert len(out["exclude"]) == 2
    auto_ops, _ = prep.plan_create_grid_titles(
        analysis, {"findings": []}, None,
        exclude=out["exclude"], reserved=out["reserved"],
        pre_titled=out["titled_cells"], pre_inserted=out["inserted_rows"],
    )
    assert not [o for o in auto_ops if o["sheet"] == "S"], auto_ops


def test_retitling_releases_the_old_name_for_another_area(tmp_path):
    analysis = _analysis(tmp_path, {
        "S": {"B2": "#Old /Input", "B3": "H1", "C3": "H2", "B4": 1, "C4": 2, "E3": "X", "F3": "Y", "E4": 1, "F4": 2},
    })
    grids = {g["ref"]: g for g in prep.all_grids(analysis)}
    assert grids["B3:C4"]["name"] == "Old"
    out = prep.plan_named_areas(analysis, [
        {"sheet": "S", "ref": "B3", "name": "New", "flags": []},
        {"sheet": "S", "ref": "E3:F4", "name": "Old", "flags": []},
    ])
    titles = {o["cell"]: o["after"] for o in out["ops"]}
    assert titles["B2"] == "#New"
    assert titles["E2"] == "#Old"  # not "#Old (2)": the rename released it


def test_unknown_flag_refuses_the_area(tmp_path):
    analysis = _analysis(tmp_path, {"Data": {"B3": "H", "C3": "K", "B4": 1, "C4": 2}})
    out = prep.plan_named_areas(analysis, [{"sheet": "Data", "ref": "B3", "name": "X", "flags": ["not a flag"]}])
    assert not out["ops"]
    assert any("unrecognisable flag" in s for s in out["skipped"]), out["skipped"]


def test_a_title_cell_something_reads_is_not_rewritten(tmp_path):
    analysis = _analysis(tmp_path, {
        "S": {"B2": "#Old", "B3": "H1", "C3": "H2", "B4": 1, "C4": 2},
        "T": {"A1": "=S!B2"},
    })
    out = prep.plan_named_areas(analysis, [{"sheet": "S", "ref": "B3:C4", "name": "New", "flags": []}])
    assert not out["ops"]
    assert any("title cell B2 is read by T!A1" in s for s in out["skipped"]), out["skipped"]
    assert "S!B3:C4" in out["exclude"]  # the automatic pass must not retry the same refused write


def test_automatic_pass_respects_manual_names_and_grids(tmp_path):
    """The auto titler skips the grid the user named, and a name the user took
    is not written twice: the heuristic name gets a ' (2)'."""
    analysis = _analysis(tmp_path, {
        "P": {"B3": "H", "C3": "K", "B4": 1, "C4": 2, "F3": "H2", "G3": "K2", "F4": 1, "G4": 2},
    })
    manual = prep.plan_named_areas(analysis, [{"sheet": "P", "ref": "B3", "name": "H2 K2", "flags": []}])
    assert manual["ops"][0]["cell"] == "B2" and manual["ops"][0]["after"] == "#H2 K2"
    auto_ops, _ = prep.plan_create_grid_titles(
        analysis, {"findings": []}, None,
        exclude=manual["exclude"], reserved=manual["reserved"],
        pre_titled=manual["titled_cells"], pre_inserted=manual["inserted_rows"],
    )
    assert all(o.get("cell") != "B2" for o in auto_ops), auto_ops  # the user's grid is left alone
    f2 = next(o for o in auto_ops if o.get("cell") == "F2")
    assert f2["after"] == "#H2 K2 (2)"  # the user's name is reserved


def test_two_areas_sharing_row_one_share_a_single_row_insert(tmp_path):
    analysis = _analysis(tmp_path, {"Q": {"A1": "a", "A2": 1, "C1": "b", "C2": 2}})
    out = prep.plan_named_areas(analysis, [
        {"sheet": "Q", "ref": "A1:A2", "name": "Left", "flags": []},
        {"sheet": "Q", "ref": "C1:C2", "name": "Right", "flags": []},
    ])
    inserts = [o for o in out["ops"] if o["op"] == "insert_row"]
    writes = [o for o in out["ops"] if o["op"] == "set_value"]
    assert len(inserts) == 1 and inserts[0]["row"] == 1
    assert {w["cell"] for w in writes} == {"A1", "C1"} and all(w.get("after_inserts") for w in writes)


# --- the endpoints behind the Grid Namer screen ----------------------------------------
@pytest.fixture(scope="module")
def client():
    return TestClient(server.app)


def _upload(client, path: Path) -> dict:
    with path.open("rb") as f:
        res = client.post("/api/sessions", files={"file": (path.name, f, "application/octet-stream")}, data={"mode": "plan"})
    assert res.status_code == 200, res.text
    return res.json()


def test_mind_flags_endpoint_lists_documented_title_flags(client):
    flags = client.get("/api/mind-flags").json()["flags"]
    names = {f["name"] for f in flags}
    assert "Input" in names and "Export" in names
    assert all(f.get("meaning") for f in flags if f["name"] == "Input")


def test_sheet_cells_and_grid_namer_endpoints(client, tmp_path):
    path = _wb(tmp_path / "namer.xlsx", {"Model": {"B3": "Age", "C3": "Rate", "B4": 30, "C4": 0.01, "E7": "=C4*2"}})
    body = _upload(client, path)
    sid = body["sessionId"]

    cells = client.get(f"/api/sessions/{sid}/sheet-cells", params={"sheet": "Model"}).json()
    assert cells["rows"][2][1] == "Age" and cells["rows"][2][2] == "Rate"
    assert cells["rows"][6][4] == "=C4*2"  # no cached value in a synthetic file -> the formula text
    assert cells["truncated"] is False and cells["n_rows"] == 7
    assert client.get(f"/api/sessions/{sid}/sheet-cells", params={"sheet": "Nope"}).status_code == 404

    res = client.post(f"/api/sessions/{sid}/grid-namer", json={
        "areas": [{"sheet": "Model", "ref": "B3:C4", "name": "Mortality", "flags": ["input"]}],
        "nameRest": True,
    })
    assert res.status_code == 200, res.text
    out = res.json()
    assert out["result"]["status"] == "APPLIED" and out["version"]["source"] == "grid_namer"
    assert out["named"] == ["Model!B3:C4"] and out["manual_ops"] == 1
    grid = next(g for s in out["summary"]["sheets"] for g in s["grids"] if g["sheet"] == "Model" and "B" in g["ref"])
    assert grid["name"] == "Mortality" and grid["flag_names"] == ["input"]

    bad = client.post(f"/api/sessions/{sid}/grid-namer", json={"areas": [{"sheet": "Model", "ref": "ZZ99", "name": "Nothing"}]})
    assert bad.status_code == 422 and "no area could be named" in bad.json()["detail"]
