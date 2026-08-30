import pytest

from app.config import load_config
from app.inventory import build_analysis
from app.rules_engine import RulesEngine
from app.validators import run_rule


@pytest.fixture(scope="module")
def engine():
    return RulesEngine()


@pytest.fixture(scope="module")
def config():
    return load_config()


def _finding(engine, config, analysis, rule_id):
    rule = engine.get(rule_id)
    assert rule is not None, f"{rule_id} missing from loaded rules"
    return run_rule(rule, analysis, config)


def _statuses(engine, config, analysis, rule_ids):
    return {rid: _finding(engine, config, analysis, rid)["status"] for rid in rule_ids}


# --- loops (1.1.0 behaviour preserved) ------------------------------------------
def test_loop_inventory_and_naming_pass_on_clean_fixture(tmp_path, engine, config, loops_ok_xlsx):
    analysis = build_analysis(loops_ok_xlsx, tmp_path / "work", "t1")
    assert _statuses(engine, config, analysis, ["LOOP-001", "LOOP-002", "LOOP-003", "LBL-001"]) == {
        "LOOP-001": "PASS", "LOOP-002": "PASS", "LOOP-003": "PASS", "LBL-001": "PASS"}


def test_loop_validators_fail_on_broken_fixture(tmp_path, engine, config, loops_broken_xlsx):
    analysis = build_analysis(loops_broken_xlsx, tmp_path / "work", "t2")
    naming = _finding(engine, config, analysis, "LOOP-002")
    assert naming["status"] == "ERROR"
    assert "scenario" in str(naming["observed"]).lower()
    assert _finding(engine, config, analysis, "LOOP-003")["status"] == "ERROR"
    assert _finding(engine, config, analysis, "LBL-001")["status"] == "ERROR"


# --- formulas -----------------------------------------------------------------------
def test_unsupported_functions_flags_unregistered_mm_call(tmp_path, engine, config, unregistered_function_xlsx):
    analysis = build_analysis(unregistered_function_xlsx, tmp_path / "work", "t3")
    finding = _finding(engine, config, analysis, "FRM-002")
    assert finding["status"] == "ERROR"
    assert "MM_TOTALLYMADEUP" in finding["observed"]["unregistered_mm_functions"]


def test_unsupported_functions_passes_on_registered_calls(tmp_path, engine, config, registered_functions_only_xlsx):
    analysis = build_analysis(registered_functions_only_xlsx, tmp_path / "work", "t4")
    assert _finding(engine, config, analysis, "FRM-002")["status"] == "PASS"


def test_looplabels_is_registered_via_the_kb_registry(tmp_path, engine, config, loops_ok_xlsx):
    """1.2.0 flagged MM_LOOPLABELS because doc 02 has no entry for it; the KB
    registry mined from CompleteMindDocn.docx (1.3.0) documents it."""
    analysis = build_analysis(loops_ok_xlsx, tmp_path / "work", "t4b")
    finding = _finding(engine, config, analysis, "FRM-002")
    assert finding["status"] == "PASS"
    assert finding["observed"]["unregistered_mm_functions"] == []


def test_native_function_off_the_kb_list_and_vba_udf_are_flagged(tmp_path, engine, config, array_formula_xlsx):
    analysis = build_analysis(array_formula_xlsx, tmp_path / "work", "t4c")
    frm = _finding(engine, config, analysis, "FRM-002")
    assert frm["status"] == "ERROR"
    assert "XLOOKUP" in frm["observed"]["unsupported_native_functions"]
    assert "MYMACROFUNCTION" not in frm["observed"]["unsupported_native_functions"]
    udf = _finding(engine, config, analysis, "FORMULA-002")
    assert udf["status"] == "WARNING"  # no VBA project in this fixture -> unknown function, not a confirmed UDF
    assert "MYMACROFUNCTION" in udf["observed"]["udf_candidates"]
    dyn = _finding(engine, config, analysis, "FRM-003")
    assert dyn["status"] == "WARNING"
    assert dyn["observed"]["counts"]["implicit_intersection"] == 1
    assert dyn["observed"]["counts"]["array_formulas"] == 1


def test_array_formula_past_grid_boundary_is_oversized(tmp_path, engine, config, array_formula_xlsx):
    analysis = build_analysis(array_formula_xlsx, tmp_path / "work", "t4d")
    finding = _finding(engine, config, analysis, "RSK-002")
    assert finding["status"] == "ERROR"
    assert finding["observed"]["oversized_array_formulas"][0]["cell"] == "C5"


def test_vba_presence(tmp_path, engine, config, plain_grid_xlsx, macro_xlsm):
    clean = build_analysis(plain_grid_xlsx, tmp_path / "work1", "t5")
    assert _finding(engine, config, clean, "FRM-004")["status"] == "PASS"
    with_vba = build_analysis(macro_xlsm, tmp_path / "work2", "t6")
    # Mind ignores VBA it can't run -- presence alone isn't a blocker, only WARNING.
    assert _finding(engine, config, with_vba, "FRM-004")["status"] == "WARNING"


# --- format / structure / risk ------------------------------------------------------
def test_format_and_structure_inventory_warn_on_messy_fixture(tmp_path, engine, config, messy_workbook_xlsx):
    analysis = build_analysis(messy_workbook_xlsx, tmp_path / "work", "t7")
    assert _statuses(engine, config, analysis, ["FMT-004", "FMT-005", "FMT-007", "RSK-003", "STR-007"]) == {
        "FMT-004": "WARNING", "FMT-005": "WARNING", "FMT-007": "WARNING", "RSK-003": "WARNING", "STR-007": "WARNING"}


def test_clean_fixture_passes_format_and_risk_checks(tmp_path, engine, config, plain_grid_xlsx):
    analysis = build_analysis(plain_grid_xlsx, tmp_path / "work", "t8")
    assert _statuses(engine, config, analysis, ["FMT-001", "FMT-002", "FMT-003", "FMT-004", "FMT-006", "RSK-003", "RSK-004", "MMX-002"]) == {
        "FMT-001": "PASS", "FMT-002": "PASS", "FMT-003": "PASS", "FMT-004": "PASS", "FMT-006": "PASS", "RSK-003": "PASS", "RSK-004": "PASS", "MMX-002": "PASS"}


def test_recalculation_always_not_supported(tmp_path, engine, config, plain_grid_xlsx):
    analysis = build_analysis(plain_grid_xlsx, tmp_path / "work", "t9")
    assert _finding(engine, config, analysis, "READY-001")["status"] == "NOT_SUPPORTED"


def test_every_active_rule_has_a_real_validator(engine):
    """1.3.0: no active rule falls through to the 'Not implemented' path any more."""
    missing = [r["id"] for r in engine.active_rules() if RulesEngine.resolve_validator(r["validation"]["implementation"]) is None]
    assert missing == []


def test_unresolvable_implementation_still_reports_on_unevaluable(tmp_path, engine, config, plain_grid_xlsx):
    """The honest fallback (SKILL.md rule #8) is still wired for any future rule."""
    analysis = build_analysis(plain_grid_xlsx, tmp_path / "work", "t10")
    rule = dict(engine.get("RES-001"))
    rule["validation"] = {**rule["validation"], "implementation": "validators.loop.does_not_exist"}
    finding = run_rule(rule, analysis, config)
    assert finding["status"] == rule["on_unevaluable"]
    assert "Not implemented" in finding["message"]


# --- the KB-backed validators on a well-formed model -----------------------------------
GOOD_EXPECTATIONS = {
    "STR-001": "PASS", "STR-003": "PASS", "STR-004": "PASS", "STR-005": "PASS", "STR-006": "PASS",
    "INP-001": "PASS", "INP-002": "PASS", "INP-003": "PASS", "INP-004": "PASS", "INP-005": "PASS", "INP-006": "PASS",
    "EXP-001": "PASS", "EXP-002": "PASS", "EXP-003": "PASS", "EXP-004": "PASS",
    "PAR-001": "PASS", "PAR-002": "PASS", "PAR-003": "PASS", "PAR-004": "PASS",
    "PRJ-001": "PASS", "PRJ-002": "PASS", "PRJ-003": "WARNING", "PRJ-005": "PASS", "PRJ-006": "PASS",
    "RES-001": "PASS", "RES-002": "PASS", "RES-004": "PASS", "RES-005": "PASS",
    "RZS-001": "PASS", "RZS-002": "PASS", "RZS-003": "PASS", "RZS-004": "PASS", "RZS-005": "PASS", "RZS-007": "PASS",
    "RSK-001": "PASS", "RSK-002": "PASS",
    "LKP-001": "PASS", "LKP-002": "PASS", "LKP-003": "PASS", "LKP-004": "PASS", "LKP-005": "PASS",
    "CAL-001": "PASS", "CAL-002": "PASS", "CAL-004": "PASS", "CAL-005": "PASS",
    "MMX-001": "PASS", "MMX-002": "PASS", "MMX-003": "PASS",
    "FLG-001": "PASS", "GRP-001": "PASS", "BKP-001": "PASS", "HID-001": "PASS", "RNG-001": "PASS", "UNQ-001": "PASS", "TRN-001": "PASS", "SIM-001": "PASS",
    "DBG-003": "PASS", "DBG-004": "PASS",
}


def test_well_formed_model_passes_the_kb_backed_validators(tmp_path, engine, config, flagged_model_xlsx):
    analysis = build_analysis(flagged_model_xlsx, tmp_path / "work", "good")
    actual = {}
    for rid, expected in GOOD_EXPECTATIONS.items():
        f = _finding(engine, config, analysis, rid)
        actual[rid] = (f["status"], f["message"][:120])
    mismatches = {rid: actual[rid] for rid, expected in GOOD_EXPECTATIONS.items() if actual[rid][0] != expected}
    assert mismatches == {}


BROKEN_EXPECTATIONS = {
    "STR-001": "WARNING",  # '#Second' trapped inside the first grid
    "STR-003": "ERROR",    # mixed header row (Key, Key, 42) on a grid whose columns Mind matches by header
    "INP-005": "ERROR",    # /Reorder without /Input, duplicate headers
    "FLG-001": "WARNING",  # /Inpt is not a documented flag
    "PAR-002": "ERROR",    # 'Kind' instead of 'Type'
    "PRJ-001": "ERROR",    # two /ProjectSettings grids
    "PRJ-006": "ERROR",    # EnableDebugMode: a setting name Mind rejects
    "PRJ-002": "ERROR",    # 'Setting' header, 'sometimes' as Locked
    "UNQ-001": "ERROR",
    "EXP-003": "ERROR",    # unknown column, non-export GridName, bad boolean
    "EXP-004": "WARNING",  # {Bogus} tag
    "RES-002": "ERROR",    # "scenario" vs "Scenario"
    "RES-004": "ERROR",    # same loop twice
    "RZS-002": "ERROR",    # MM_SETSIZE nested in SUM
    "RZS-003": "ERROR",    # 'leftover' in the destination cells
    "RZS-005": "WARNING",  # single-member resize group
    "GRP-001": "ERROR",    # x is not an integer
    "BKP-001": "ERROR",    # no destination, button inside the source grid
    "HID-001": "ERROR",    # 'maybe'
    "SIM-001": "ERROR",    # 'lots'
    "TRN-001": "ERROR",    # 'english'
    "RNG-001": "ERROR",    # MM_RANGE shares its grid with AA5
    "LKP-002": "ERROR",    # range starts on a data row
    "LKP-003": "ERROR",    # 'Result' is not a header
    "LKP-005": "WARNING",  # plain MM_READTABLE
}


def test_broken_model_fails_each_kb_backed_validator(tmp_path, engine, config, flagged_model_broken_xlsx):
    analysis = build_analysis(flagged_model_broken_xlsx, tmp_path / "work", "bad")
    actual = {}
    for rid, expected in BROKEN_EXPECTATIONS.items():
        f = _finding(engine, config, analysis, rid)
        actual[rid] = (f["status"], f["message"][:160])
    mismatches = {rid: actual[rid] for rid, expected in BROKEN_EXPECTATIONS.items() if actual[rid][0] != expected}
    assert mismatches == {}


def test_par_003_and_par_004_report_bad_types_and_ranges(tmp_path, engine, config, flagged_model_broken_xlsx):
    analysis = build_analysis(flagged_model_broken_xlsx, tmp_path / "work", "par")
    # PAR-002 already fails on the header; PAR-003/004 still inspect the rows they can.
    t = _finding(engine, config, analysis, "PAR-003")
    assert t["status"] == "ERROR" and "numbre" in str(t["observed"])
    p = _finding(engine, config, analysis, "PAR-004")
    assert p["status"] == "ERROR"


def test_ref_001_flags_broken_defined_name_references(tmp_path, engine, config):
    """REF-001: a cell referencing a #REF! defined name is an ERROR blocker."""
    import openpyxl
    from openpyxl.workbook.defined_name import DefinedName

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"] = "#Data"
    ws["A2"] = "x"
    ws["A3"] = "=COLUMNS(NB_Breakdown)"
    dn = DefinedName("NB_Breakdown", attr_text="#REF!")
    try:
        wb.defined_names["NB_Breakdown"] = dn
    except Exception:
        wb.defined_names.add(dn)
    p = tmp_path / "broken_name.xlsx"
    wb.save(p)
    wb.close()
    analysis = build_analysis(p, tmp_path / "w", "b")
    f = _finding(engine, config, analysis, "REF-001")
    assert f["status"] == "ERROR"
    assert "NB_Breakdown" in str(f["observed"]["broken_names"])
    assert any(s["cell"] == "A3" for s in f["observed"]["sites"])


def test_filter_is_treated_as_mind_supported():
    """FILTER is a native function the KB scrape omits but real Mind conversion accepts
    (proven by uploading the Shlomo model); FRM-002 must not flag it as unsupported."""
    from app.validators.formula import _supported_native_functions, MIND_CONFIRMED_SUPPORTED

    supported = _supported_native_functions()
    assert "FILTER" in supported
    assert MIND_CONFIRMED_SUPPORTED <= supported
    assert len(supported) > 100  # the KB list still loads


def test_rep_001_parse_total_recognises_pure_totals():
    from app.validators.structure import _parse_total

    assert _parse_total("=SUM(A1:A10)") == ("sum", "A1:A10")
    assert _parse_total("=SUM($B$2:$B$5)") == ("sum", "B2:B5")
    assert _parse_total("=A1+A2+A3") == ("chain", ["A1", "A2", "A3"])
    # not pure totals
    assert _parse_total("=SUM(A1:A10)+B1") is None
    assert _parse_total("=SUM(Sheet2!A1:A10)") is None
    assert _parse_total("=A1*2") is None
    assert _parse_total("=A1") is None
