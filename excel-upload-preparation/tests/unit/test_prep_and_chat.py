"""1.4.0: the prep planner/executor (app/prep.py) and the grounded chatbot
(app/chat_context.py). LLM calls are mocked; one Excel-backed test proves the
structural operations (rename / insert row) really go through Excel."""
from pathlib import Path

import pytest

from app import chat_context, llm
from app.config import load_config
from app.excel_com import com_available
from app.inventory import build_analysis
from app.modes import plan_mode
from app.prep import apply_operations, formula_replacement_op, plan_actions
from app.rules_engine import RulesEngine
from app.validators import run_rule

needs_excel = pytest.mark.skipif(not com_available(), reason="requires pywin32 + an installed Excel")


def _run(tmp_path, source):
    engine = RulesEngine()
    result = plan_mode.run(source, tmp_path / "work", load_config(), engine=engine)
    return result, engine


def _by_id(plan):
    return {a["id"]: a for a in plan}


def test_plan_on_broken_model_lists_concrete_operations_and_skips(tmp_path, flagged_model_broken_xlsx):
    result, _ = _run(tmp_path, flagged_model_broken_xlsx)
    plan = _by_id(plan_actions(result["workbook_analysis"], result["validation_report"]))
    # /Inpt -> /Input (FLG-001) and /Reorder gets /Input (INP-005): same title cell, both planned
    assert plan["flag_spelling"]["count"] == 1 and "/Input" in plan["flag_spelling"]["operations"][0]["after"]
    assert plan["reorder_input_flag"]["count"] == 1
    # "scenario" -> "Scenario" in the MM_RESULT (RES-002 case mismatch)
    ops = plan["loop_name_case"]["operations"]
    assert any(o["cell"] == "J6" and '"Scenario"' in o["after"] for o in ops)
    # 'Kind' -> 'Type' (PAR-002), 'Setting' -> 'Name' (PRJ-002) header corrections
    headers = {(o["before"], o["after"]) for o in plan["special_headers"]["operations"]}
    assert ("Kind", "Type") in headers and ("Setting", "Name") in headers
    # the trapped '#Second' title shares rows with other grids -> skipped with a reason, never a blind row insert
    assert plan["separate_merged_grids"]["count"] == 0
    assert any("split" in s for s in plan["separate_merged_grids"]["skipped"])
    # plans never touch the file
    assert flagged_model_broken_xlsx.stat().st_size > 0


def test_plan_on_clean_fixture_is_empty(tmp_path, plain_grid_xlsx):
    result, _ = _run(tmp_path, plain_grid_xlsx)
    plan = plan_actions(result["workbook_analysis"], result["validation_report"])
    # the plain grid has no '#' title, so only the title action has work to do
    assert all(a["count"] == 0 for a in plan if a["id"] != "create_grid_titles")
    assert {a["id"]: a["count"] for a in plan}["create_grid_titles"] >= 1


def test_apply_value_and_formula_operations_without_excel_improves_findings(tmp_path, flagged_model_broken_xlsx):
    result, engine = _run(tmp_path, flagged_model_broken_xlsx)
    plan = _by_id(plan_actions(result["workbook_analysis"], result["validation_report"]))
    ops = plan["flag_spelling"]["operations"] + plan["loop_name_case"]["operations"] + plan["special_headers"]["operations"]
    before = flagged_model_broken_xlsx.read_bytes()
    out = apply_operations(flagged_model_broken_xlsx, tmp_path / "prep", ops, prefer_excel=False)
    assert out["status"] == "APPLIED" and out["method"] == "openpyxl"
    assert flagged_model_broken_xlsx.read_bytes() == before  # source untouched
    log = Path(str(out["output_path"]) + ".changelog.json")
    assert log.is_file()
    after = build_analysis(Path(out["output_path"]), tmp_path / "after", "a")
    cfg = load_config()
    assert run_rule(engine.get("FLG-001"), after, cfg)["status"] == "PASS"
    assert run_rule(engine.get("RES-002"), after, cfg)["status"] == "PASS"
    assert run_rule(engine.get("PAR-002"), after, cfg)["status"] == "PASS"


def test_structural_operations_are_refused_without_excel(tmp_path, messy_workbook_xlsx):
    result, _ = _run(tmp_path, messy_workbook_xlsx)
    plan = _by_id(plan_actions(result["workbook_analysis"], result["validation_report"]))
    assert plan["hide_marker"]["count"] == 1 and plan["hide_marker"]["operations"][0]["after"] == "Notes &&Hide"
    assert plan["unprotect_sheets"]["count"] == 1
    out = apply_operations(messy_workbook_xlsx, tmp_path / "prep", plan["hide_marker"]["operations"] + plan["unprotect_sheets"]["operations"], prefer_excel=False)
    assert out["status"] == "ERROR" and len(out["failed"]) == 2
    assert all("requires Excel" in f["error"] for f in out["failed"])


def test_formula_replacement_op_requires_a_formula():
    with pytest.raises(ValueError):
        formula_replacement_op("S", "A1", "=XLOOKUP(1,A:A,B:B)", "hello", "FRM-002")
    op = formula_replacement_op("S", "A1", "=XLOOKUP(1,A:A,B:B)", " =INDEX(B:B,MATCH(1,A:A,0)) ", "FRM-002")
    assert op["op"] == "set_formula" and op["after"] == "=INDEX(B:B,MATCH(1,A:A,0))"


@needs_excel
def test_excel_applies_rename_insert_and_unprotect_and_output_opens(tmp_path, messy_workbook_xlsx):
    result, engine = _run(tmp_path, messy_workbook_xlsx)
    plan = _by_id(plan_actions(result["workbook_analysis"], result["validation_report"]))
    ops = plan["hide_marker"]["operations"] + plan["unprotect_sheets"]["operations"]
    out = apply_operations(messy_workbook_xlsx, tmp_path / "prep", ops)
    assert out["status"] == "APPLIED" and out["method"] == "excel_com"
    assert out["verified_opens_in_excel"] is True
    after = build_analysis(Path(out["output_path"]), tmp_path / "after", "b")
    cfg = load_config()
    assert run_rule(engine.get("STR-007"), after, cfg)["status"] == "PASS"
    assert run_rule(engine.get("FMT-005"), after, cfg)["status"] == "PASS"
# --- chatbot ---------------------------------------------------------------------------


def test_context_pack_and_retrieval_are_grounded(tmp_path, flagged_model_broken_xlsx):
    result, _ = _run(tmp_path, flagged_model_broken_xlsx)
    plan = plan_actions(result["workbook_analysis"], result["validation_report"])
    ctx = chat_context.build_context_pack(result["workbook_analysis"], result["validation_report"], plan)
    assert "## Model" in ctx and "## Settings" in ctx and "/Reorder /Inpt" in ctx
    assert "[ERROR] RES-002" in ctx and "Prep actions" in ctx
    assert len(ctx) <= chat_context.MAX_CONTEXT_CHARS
    detail = chat_context.retrieve("why does RES-002 fail, and what is in Model!J6? show the Params grid", result["workbook_analysis"], result["validation_report"])
    assert "## Finding RES-002" in detail
    assert "J6 = =MM_RESULT(J5,\"scenario\",1)" in detail
    assert "## Grid Params" in detail and "PossibleValues" in detail


def test_answer_question_keeps_history_and_uses_the_gateway(tmp_path, plain_grid_xlsx, monkeypatch):
    result, _ = _run(tmp_path, plain_grid_xlsx)
    captured = {}

    def fake_chat(messages, system_prompt, **kwargs):
        captured["messages"] = messages
        captured["system"] = system_prompt
        return {"available": True, "text": "The Data sheet has one grid.", "message": None}

    monkeypatch.setattr(chat_context, "chat_completion", fake_chat)
    history = [{"role": "user", "content": "hello"}, {"role": "assistant", "content": "hi"}]
    reply = chat_context.answer_question(history, "what is on sheet Data?", result["workbook_analysis"], result["validation_report"])
    assert reply["text"] == "The Data sheet has one grid."
    assert [m["role"] for m in captured["messages"]] == ["user", "assistant", "user"]
    assert "<workbook_context>" in captured["system"] and "Data" in captured["system"]
    assert "## Sheet Data" in captured["messages"][-1]["content"]


def test_chat_completion_payload_shape(monkeypatch):
    calls = {}

    class FakeResp:
        def raise_for_status(self):
            pass

        def json(self):
            return {"content": [{"type": "text", "text": "ok"}]}

    def fake_post(url, headers=None, json=None, timeout=None):
        calls["payload"] = json
        return FakeResp()

    monkeypatch.setattr(llm, "find_api_key", lambda start_dir=None: ("key", "https://example.invalid"))
    import requests

    monkeypatch.setattr(requests, "post", fake_post)
    out = llm.chat_completion([{"role": "user", "content": "q1"}, {"role": "assistant", "content": "a1"}, {"role": "user", "content": "q2"}], "SYS")
    assert out["text"] == "ok"
    msgs = calls["payload"]["messages"]
    assert msgs[0] == {"role": "system", "content": "SYS"}
    assert [m["role"] for m in msgs[1:]] == ["user", "assistant", "user"]
    assert llm.extract_formula("Use this:\n=INDEX(A:A,1)\nbecause...") == "=INDEX(A:A,1)"
    assert llm.extract_formula("no formula here") is None
# --- context-aware grid titles (1.5.0) ------------------------------------------------


def test_create_grid_titles_uses_surrounding_context(tmp_path, untitled_grids_xlsx):
    result, _ = _run(tmp_path, untitled_grids_xlsx)
    plan = _by_id(plan_actions(result["workbook_analysis"], result["validation_report"]))
    a = plan["create_grid_titles"]
    titles = {(o["sheet"], o["cell"]): o["after"] for o in a["operations"] if o["op"] == "set_value"}
    # label above with a blank row between -> title in the empty cell above, label moved
    assert titles[("Data", "B3")] == "#Premium table"
    assert any(o["op"] == "clear_cell" and o["cell"] == "B2" for o in a["operations"])
    # caption directly above a two-column table -> caption becomes the title (one grid, not two)
    assert titles[("Data", "F4")] == "#Rates"
    # label to the left -> title, label moved
    assert titles[("Data", "L3")] == "#Factors"
    assert any(o["op"] == "clear_cell" and o["cell"] == "J4" for o in a["operations"])
    # header row only -> name from the headers
    assert titles[("Data", "O3")] == "#Year Amount"
    # grid on row 1 -> insert a row, then title written at the post-insert coordinate
    assert any(o["op"] == "insert_row" and o["sheet"] == "Top" and o["row"] == 1 for o in a["operations"])
    top_title = [o for o in a["operations"] if o["sheet"] == "Top" and o["op"] == "set_value"][0]
    assert top_title["cell"] == "A1" and top_title["after"] == "#X Y" and top_title.get("after_inserts")
    # every title target is distinct and every name unique
    targets = [(o["sheet"], o["cell"]) for o in a["operations"] if o["op"] == "set_value"]
    assert len(targets) == len(set(targets))
    names = [o["after"] for o in a["operations"] if o["op"] == "set_value"]
    assert len(names) == len(set(names))
    # the far-away standalone label is reported, not touched
    assert any("A1" in s and "not next to a grid" in s for s in a["skipped"])


@needs_excel
def test_create_grid_titles_applied_by_excel_makes_every_grid_titled(tmp_path, untitled_grids_xlsx):
    from app.grids import all_grids

    result, engine = _run(tmp_path, untitled_grids_xlsx)
    plan = _by_id(plan_actions(result["workbook_analysis"], result["validation_report"]))
    out = apply_operations(untitled_grids_xlsx, tmp_path / "prep", plan["create_grid_titles"]["operations"])
    assert out["status"] == "APPLIED" and out["verified_opens_in_excel"] is True
    after = build_analysis(Path(out["output_path"]), tmp_path / "after", "t")
    grids = all_grids(after)
    assert all(g["name"] for g in grids), [g["display_name"] for g in grids if not g["name"]]
    assert sorted(g["name"] for g in grids if g["sheet"] == "Data") == ["Factors", "Premium table", "Rates", "Year Amount"]
    rates = next(g for g in grids if g["name"] == "Rates")
    assert rates["ref"] == "F5:G7"  # one grid, not a one-column block plus a leftover
    assert run_rule(engine.get("STR-004"), after, load_config())["status"] == "PASS"
# --- assistant-proposed changes (1.5.0) -----------------------------------------------


def test_validate_proposal_accepts_documented_ops_and_rejects_bad_ones(tmp_path, flagged_model_broken_xlsx):
    from app.prep import validate_proposal

    result, _ = _run(tmp_path, flagged_model_broken_xlsx)
    proposal = {"summary": "tidy", "operations": [
        {"op": "set_value", "sheet": "model", "cell": "A3", "value": "#Assumptions /Input /Reorder"},
        {"op": "set_formula", "sheet": "Model", "cell": "J6:J7", "formula": "=MM_RESULT(J5,\"Scenario\",1)"},
        {"op": "rename_sheet", "sheet": "Settings", "new_name": "Config"},
        {"op": "insert_row", "sheet": "Model", "row": 3},
        {"op": "set_value", "sheet": "Nowhere", "cell": "A1", "value": 1},
        {"op": "set_value", "sheet": "Model", "cell": "A:A", "value": 1},
        {"op": "set_value", "sheet": "Model", "cell": "A1", "value": "=1+1"},
        {"op": "drop_sheet", "sheet": "Model"},
    ]}
    ops, errors = validate_proposal(proposal, result["workbook_analysis"])
    assert [o["op"] for o in ops] == ["set_value", "set_formula", "rename_sheet", "insert_row"]
    assert ops[0]["sheet"] == "Model" and ops[0]["before"] == "#Assumptions /Reorder /Inpt"
    assert ops[1]["cell"] == "J6:J7" and ops[1]["before"] == "2 cells"
    assert len(errors) == 4


def test_answer_question_returns_validated_proposal_and_does_lookups(tmp_path, flagged_model_broken_xlsx, monkeypatch):
    result, _ = _run(tmp_path, flagged_model_broken_xlsx)
    replies = iter([
        '```lookup\n{"cells": ["Model!A3:C4"], "sheets": ["Settings"]}\n```',
        'I will rename the sheet.\n```changes\n{"summary": "rename Settings", "operations": [{"op": "rename_sheet", "sheet": "Settings", "new_name": "Config"}, {"op": "set_value", "sheet": "Model", "cell": "ZZ9", "value": 1}]}\n```',
    ])
    seen = []

    def fake_chat(messages, system_prompt, **kwargs):
        seen.append([m["content"] for m in messages])
        return {"available": True, "text": next(replies), "message": None}

    monkeypatch.setattr(chat_context, "chat_completion", fake_chat)
    reply = chat_context.answer_question([], "rename the Settings sheet to Config", result["workbook_analysis"], result["validation_report"])
    assert reply["lookups"] == 1
    assert "<lookup_result>" in seen[1][-1] and "A3 = #Assumptions /Reorder /Inpt" in seen[1][-1] and "## Sheet Settings" in seen[1][-1]
    assert reply["text"] == "I will rename the sheet."
    assert [o["op"] for o in reply["proposal"]] == ["rename_sheet", "set_value"]
    assert reply["proposal"][0]["after"] == "Config" and reply["proposal_summary"] == "rename Settings"
    assert reply["proposal_errors"] == []


def test_apply_assistant_proposal_without_excel(tmp_path, flagged_model_broken_xlsx):
    from app.prep import validate_proposal

    result, engine = _run(tmp_path, flagged_model_broken_xlsx)
    ops, errors = validate_proposal({"summary": "fix title", "operations": [{"op": "set_value", "sheet": "Model", "cell": "A3", "value": "#Assumptions /Input /Reorder"}]}, result["workbook_analysis"])
    assert not errors
    out = apply_operations(flagged_model_broken_xlsx, tmp_path / "chat", ops, prefer_excel=False)
    assert out["status"] == "APPLIED"
    after = build_analysis(Path(out["output_path"]), tmp_path / "after", "c")
    assert run_rule(engine.get("FLG-001"), after, load_config())["status"] == "PASS"  # the misspelt /Inpt is gone
    grid = next(g for g in after["workbooks"][0]["sheets"][0]["grids"] if g["anchor"] == "A4")
    assert grid["flag_names"] == ["input", "reorder"]
# --- re-applying to an output (1.5.1 regression) ---------------------------------------


def test_apply_twice_to_the_same_work_dir_never_copies_a_file_onto_itself(tmp_path, flagged_model_broken_xlsx):
    """The UI re-applies changes to the file it just produced (same work_dir);
    the copy helper must never target the source path itself."""
    result, engine = _run(tmp_path, flagged_model_broken_xlsx)
    plan = _by_id(plan_actions(result["workbook_analysis"], result["validation_report"]))
    first = apply_operations(flagged_model_broken_xlsx, tmp_path / "prep", plan["flag_spelling"]["operations"], prefer_excel=False)
    assert first["status"] == "APPLIED"
    produced = Path(first["output_path"])
    second = apply_operations(produced, tmp_path / "prep", plan["special_headers"]["operations"], prefer_excel=False)
    assert second["status"] == "APPLIED"
    assert Path(second["output_path"]) != produced and Path(second["output_path"]).is_file()
    assert produced.is_file()  # the first output is left intact
    after = build_analysis(Path(second["output_path"]), tmp_path / "after", "twice")
    cfg = load_config()
    assert run_rule(engine.get("FLG-001"), after, cfg)["status"] == "PASS"
    assert run_rule(engine.get("PAR-002"), after, cfg)["status"] == "PASS"
# --- 1.6.2 regression: recalculate -> apply through Excel -> recalculate again ------------------


@pytest.mark.skipif(not com_available(), reason="requires pywin32 + an installed Excel")
def test_recalculation_still_runs_after_an_excel_apply(tmp_path, array_formula_xlsx):
    """The hand-rolled COM lifecycle recalc.py used to have released proxies after
    CoUninitialize; the next recalculation in the same process then failed with
    'The interface is unknown'. Same-process sequence the Fix panel performs."""
    from app.inventory import make_immutable_copy
    from app.recalc import recalculate

    first = recalculate(make_immutable_copy(array_formula_xlsx, tmp_path / "r1")[0])
    assert first["ran"] is True and first["status"] == "ERROR"
    assert any(e["cell"] == "E7" and e["error"] == "#NAME?" for e in first["formula_errors"])
    out = apply_operations(array_formula_xlsx, tmp_path / "apply", [{"op": "set_formula", "action_id": "assistant", "rule_id": "CHAT", "sheet": "Arr", "cell": "E7", "before": "=MyMacroFunction(A5)", "after": "=A5", "note": "t"}])
    assert out["status"] == "APPLIED" and out["method"] == "excel_com"
    second = recalculate(make_immutable_copy(Path(out["output_path"]), tmp_path / "r2")[0])
    assert second["ran"] is True, second["message"]
    assert not any(e["cell"] == "E7" for e in second["formula_errors"])
# --- 1.6.3: root-cause grouping of recalculation errors -------------------------------------


def test_group_errors_clusters_by_root_cause():
    from app.recalc import formula_signature, group_errors

    assert formula_signature("=A5:A6*2") == formula_signature("=B7:B9*3") == "=REF*N"
    assert formula_signature('=IF(A1="x", 1, Sheet2!B2)') == formula_signature('=IF(C9="yy", 22, Other!Z1)')
    errors = [
        {"sheet": "Arr", "cell": "C7", "error": "#N/A", "formula": "=A5:A6*2"},
        {"sheet": "Arr", "cell": "C8", "error": "#N/A", "formula": "=A5:A6*2"},
        {"sheet": "Arr", "cell": "C9", "error": "#N/A", "formula": "=A5:A6*2"},
        {"sheet": "Arr", "cell": "E7", "error": "#NAME?", "formula": "=MyMacroFunction(A5)"},
        {"sheet": "Arr", "cell": "E9", "error": "#NAME?", "formula": "=MyMacroFunction(B1)+1"},
        {"sheet": "Arr", "cell": "F1", "error": "#DIV/0!", "formula": "=A1/A2"},
    ]
    gaps = [
        {"sheet": "M", "cell": "J5", "formula": '=MM_LOOP("Scenario", A5:A7)'},
        {"sheet": "M", "cell": "J6", "formula": '=MM_LOOP("Other", B1:B3)'},
        {"sheet": "M", "cell": "K1", "formula": "=MM_RESULT(J5)"},
    ]
    groups = group_errors(errors, gaps)
    by_cause = {g["signature"]: g for g in groups}
    assert by_cause["sig:=REF*N"]["count"] == 3 and by_cause["sig:=REF*N"]["kind"] == "formula"
    assert by_cause["fn:MYMACROFUNCTION"]["count"] == 2 and "MYMACROFUNCTION" in by_cause["fn:MYMACROFUNCTION"]["cause"]
    assert by_cause["sig:=REF/REF"]["count"] == 1
    assert by_cause["mm:MM_LOOP"]["count"] == 2 and by_cause["mm:MM_LOOP"]["kind"] == "addin_gap" and "MMForExcel" in by_cause["mm:MM_LOOP"]["cause"]
    assert by_cause["mm:MM_RESULT"]["count"] == 1
    assert [g["count"] for g in groups] == sorted([g["count"] for g in groups], reverse=True)  # biggest first
    assert all(g["id"].startswith("g") for g in groups)


@pytest.mark.skipif(not com_available(), reason="needs Excel")
def test_array_formulas_are_changed_as_a_unit(array_formula_xlsx, tmp_path):
    """Excel refuses to change part of a Ctrl+Shift+Enter array: the executor
    dismantles the array when the fix covers every cell, fails clearly when
    it does not, and set_array_formula writes a whole array."""
    import shutil

    import openpyxl

    from app.prep import apply_operations, validate_proposal
    from app.recalc import group_errors, recalculate

    copy = tmp_path / "rc.xlsx"
    shutil.copy(array_formula_xlsx, copy)
    rc = recalculate(copy)
    assert rc["ran"] is True
    # the fixture's array formula {=A5:A6*2} spans C5:C9; C7:C9 overflow to #N/A
    arr_cells = {e["cell"]: e for e in rc["formula_errors"] if e["cell"] in ("C7", "C8", "C9")}
    assert len(arr_cells) == 3 and all(e.get("array") == "C5:C9" for e in arr_cells.values())
    groups = group_errors(rc["formula_errors"], rc["addin_gap_errors"])
    g = next(g for g in groups if g["count"] == 3)
    assert g["arrays"] == ["Arr!C5:C9"] and "array formula" in g["cause"] and g["cells"][0]["array"] == "C5:C9"

    base = {"action_id": "assistant", "rule_id": "CHAT", "sheet": "Arr", "before": None, "note": "t"}
    # part of the array only -> refused with a helpful message, nothing half-applied
    partial = apply_operations(array_formula_xlsx, tmp_path / "partial", [{**base, "op": "clear_cell", "cell": "C9", "after": None}])
    assert partial["applied"] == [] and len(partial["failed"]) == 1 and partial["status"] == "ERROR"
    assert "part of the array formula C5:C9" in partial["failed"][0]["error"] and "set_array_formula" in partial["failed"][0]["error"]

    # every cell covered -> array dismantled, each cell gets exactly what was proposed
    ops = [
        {**base, "op": "set_formula", "cell": "C5", "after": "=INDEX(A5:A6,1)*2"},
        {**base, "op": "set_formula", "cell": "C6", "after": "=INDEX(A5:A6,2)*2"},
        {**base, "op": "clear_cell", "cell": "C7:C9", "after": None},
    ]
    full = apply_operations(array_formula_xlsx, tmp_path / "full", ops)
    assert full["failed"] == [] and len(full["applied"]) == 3 and all(a["array"] == "C5:C9" for a in full["applied"])
    ws = openpyxl.load_workbook(full["output_path"])["Arr"]
    assert ws["C5"].value == "=INDEX(A5:A6,1)*2" and ws["C6"].value == "=INDEX(A5:A6,2)*2" and all(ws[c].value is None for c in ("C7", "C8", "C9"))

    # set_array_formula: validated from the assistant's JSON, written as one array
    analysis = {"workbooks": [{"sheets": [{"name": "Arr"}]}]}
    good, errors = validate_proposal({"summary": "arr", "operations": [{"op": "set_array_formula", "sheet": "Arr", "range": "c5:c9", "formula": "=ROW(C5:C9)"}]}, analysis)
    assert errors == [] and good[0]["range"] == "C5:C9" and good[0]["cell"] == "C5:C9"
    _, bad = validate_proposal({"operations": [{"op": "set_array_formula", "sheet": "Arr", "range": "C:C", "formula": "=1"}]}, analysis)
    assert bad and "range" in bad[0]
    arr = apply_operations(array_formula_xlsx, tmp_path / "arr", good)
    assert arr["failed"] == [] and arr["applied"][0]["array"] == "C5:C9"
    c5 = openpyxl.load_workbook(arr["output_path"])["Arr"]["C5"].value
    assert getattr(c5, "ref", None) == "C5:C9" and getattr(c5, "text", None) == "=ROW(C5:C9)"


def test_validate_proposal_refuses_partial_array_changes(tmp_path, array_formula_xlsx):
    """A proposal touching only part of a Ctrl+Shift+Enter array is refused
    with the uncovered cells named; covering the whole array (or writing it
    with set_array_formula) passes."""
    from app.prep import validate_proposal

    result, _engine = _run(tmp_path, array_formula_xlsx)
    analysis = result["workbook_analysis"]  # openpyxl records the array formula {=A5:A6*2} as C5 with array_ref C5:C9
    assert any(f.get("array_ref") == "C5:C9" for f in analysis["workbooks"][0]["formulas"])
    ops, errors = validate_proposal({"operations": [{"op": "clear_cell", "sheet": "Arr", "cell": "C9"}, {"op": "set_value", "sheet": "Arr", "cell": "F1", "value": 1}]}, analysis)
    assert [o["cell"] for o in ops] == ["F1"]  # C9 is a member of the array, not an empty cell
    assert len(errors) == 1 and "part of the array formula C5:C9" in errors[0] and "still uncovered: C5, C6, C7, C8" in errors[0]
    ops, errors = validate_proposal({"operations": [{"op": "set_array_formula", "sheet": "Arr", "range": "C7:C8", "formula": "=1"}]}, analysis)
    assert ops == [] and "still uncovered: C5, C6, C9" in errors[0]
    per_cell = [{"op": "set_formula", "sheet": "Arr", "cell": c, "formula": "=1"} for c in ("C5", "C6", "C7")] + [{"op": "clear_cell", "sheet": "Arr", "cell": "C8:C9"}]
    ops, errors = validate_proposal({"operations": per_cell}, analysis)
    assert errors == [] and len(ops) == 4
    ops, errors = validate_proposal({"operations": [{"op": "set_array_formula", "sheet": "Arr", "range": "C5:C9", "formula": "=ROW(C5:C9)"}]}, analysis)
    assert errors == [] and ops[0]["range"] == "C5:C9" and ops[0]["before"] == "=A5:A6*2"
    # no-ops are refused: the same array formula again, a cell's current value, clearing an empty cell
    ops, errors = validate_proposal({"operations": [{"op": "set_array_formula", "sheet": "Arr", "range": "C5:C9", "formula": "= a5:a6 * 2"}]}, analysis)
    assert ops == [] and "re-entering it changes nothing" in errors[0]
    from app.inventory import cell_value

    a5 = cell_value(analysis, "Arr", 5, 1)
    ops, errors = validate_proposal({"operations": [{"op": "set_value", "sheet": "Arr", "cell": "A5", "value": a5}, {"op": "clear_cell", "sheet": "Arr", "cell": "ZZ99"}]}, analysis)
    assert ops == [] and "nothing would change" in errors[0] and "already empty" in errors[1]
    from app.prep import array_size_hint

    hint = array_size_hint("C5:C9", "=A5:A6*2")
    assert "5x1 cells" in hint and "A5:A6 (2x1)" in hint and "set_array_formula on C5:C6" in hint and "clear_cell C7:C9" in hint
    wide = array_size_hint("A3:E3", "=B1:D1*2")  # a 1x3 source entered over 1x5
    assert "1x5 cells" in wide and "set_array_formula on A3:C3" in wide and "clear_cell D3:E3" in wide
    assert array_size_hint("C5", "=A5:A6*2") == "" and "has 2x1 cells" in array_size_hint("C5:C6", "=SUM(1,2)")
    _, noop = validate_proposal({"operations": [{"op": "set_array_formula", "sheet": "Arr", "range": "C5:C9", "formula": "=A5:A6*2"}]}, analysis)
    assert "set_array_formula on C5:C6 with the same formula and clear_cell C7:C9" in noop[0]  # the no-op rejection carries the concrete fix
    # shrinking the array to its source values is a real fix and covers the whole array
    ops, errors = validate_proposal({"operations": [{"op": "set_array_formula", "sheet": "Arr", "range": "C5:C6", "formula": "=A5:A6*2"}, {"op": "clear_cell", "sheet": "Arr", "cell": "C7:C9"}]}, analysis)
    assert errors == [] and [o["op"] for o in ops] == ["set_array_formula", "clear_cell"]


def test_assistant_retries_a_rejected_proposal(tmp_path, array_formula_xlsx, monkeypatch):
    result, _engine = _run(tmp_path, array_formula_xlsx)
    replies = iter([
        'Clearing C9.\n```changes\n{"summary": "clear C9", "operations": [{"op": "clear_cell", "sheet": "Arr", "cell": "C9"}]}\n```',
        'Rewriting the whole array.\n```changes\n{"summary": "rewrite C5:C9", "operations": [{"op": "set_array_formula", "sheet": "Arr", "range": "C5:C9", "formula": "=ROW(C5:C9)"}]}\n```',
    ])
    seen = []

    def fake_chat(messages, system_prompt, **kwargs):
        seen.append([m["content"] for m in messages])
        return {"available": True, "text": next(replies), "message": None}

    monkeypatch.setattr(chat_context, "chat_completion", fake_chat)
    reply = chat_context.answer_question([], "fix C9", result["workbook_analysis"], result["validation_report"])
    assert len(seen) == 2 and "<proposal_rejected>" in seen[1][-1] and "part of the array formula C5:C9" in seen[1][-1]
    assert reply["proposal_retries"] == 1 and reply["proposal_errors"] == []
    assert [o["op"] for o in reply["proposal"]] == ["set_array_formula"] and reply["proposal"][0]["range"] == "C5:C9"
    assert reply["text"] == "Rewriting the whole array."


# --- section headings, insert placement, assistant naming (1.6.7) ---------------------


def test_section_heading_above_names_the_block_and_makes_room_for_the_title(tmp_path, section_heading_grids_xlsx):
    """The cell above the block is another grid (a section number + title), so
    the title action inserts a row and names the block from that heading."""
    result, _ = _run(tmp_path, section_heading_grids_xlsx)
    plan = _by_id(plan_actions(result["workbook_analysis"], result["validation_report"]))
    ops = plan["create_grid_titles"]["operations"]
    inserts = sorted(o["row"] for o in ops if o["op"] == "insert_row" and o["sheet"] == "CF")
    assert inserts == [3, 9]
    titles = {o["after"]: o for o in ops if o["op"] == "set_value"}
    assert "#Demographic Assumptions" in titles and "#Economic Assumptions" in titles
    for name in ("#Demographic Assumptions", "#Economic Assumptions"):
        assert titles[name].get("after_inserts") and titles[name].get("insert_at")
    # planned on original coordinates ...
    assert titles["#Demographic Assumptions"]["cell"] == "C3"
    assert titles["#Economic Assumptions"]["cell"] == "C9"


def test_post_insert_coordinates_shift_for_every_earlier_insert(tmp_path, section_heading_grids_xlsx):
    """... and are resolved against *all* inserts in the applied set, so the
    second title lands one row lower than it was planned."""
    from app.prep import _ordered

    result, _ = _run(tmp_path, section_heading_grids_xlsx)
    plan = _by_id(plan_actions(result["workbook_analysis"], result["validation_report"]))
    ordered = _ordered(plan["create_grid_titles"]["operations"])
    resolved = {o["after"]: o["cell"] for o in ordered if o["op"] == "set_value"}
    assert resolved["#Demographic Assumptions"] == "C3"  # no insert above it
    assert resolved["#Economic Assumptions"] == "C10"    # the insert at row 3 pushed it down
    # ordering is stable: resolving twice gives the same answer
    assert {o["after"]: o["cell"] for o in _ordered(plan["create_grid_titles"]["operations"]) if o["op"] == "set_value"} == resolved


@needs_excel
def test_section_heading_blocks_are_all_named_after_excel_applies(tmp_path, section_heading_grids_xlsx):
    from app.grids import all_grids

    result, _ = _run(tmp_path, section_heading_grids_xlsx)
    plan = _by_id(plan_actions(result["workbook_analysis"], result["validation_report"]))
    out = apply_operations(section_heading_grids_xlsx, tmp_path / "prep", plan["create_grid_titles"]["operations"])
    assert out["status"] == "APPLIED" and out["verified_opens_in_excel"] is True
    after = build_analysis(Path(out["output_path"]), tmp_path / "after", "t")
    named = {g["name"] for g in all_grids(after) if g["name"]}
    assert {"Demographic Assumptions", "Economic Assumptions"} <= named
    assert all(g["name"] for g in all_grids(after)), [g["display_name"] for g in all_grids(after) if not g["name"]]


def test_is_weak_name_spots_the_fallbacks_worth_replacing():
    from app.prep import is_weak_name

    assert is_weak_name(None, "Cashflows")
    assert is_weak_name("", "Cashflows")
    assert is_weak_name("Cashflows C4", "Cashflows")  # '<Sheet> <Anchor>' fallback
    assert is_weak_name("I", "Information")
    assert is_weak_name("I I ...", "Information")
    assert not is_weak_name("Demographic Assumptions", "Cashflows")
    assert not is_weak_name("Cashflows Reinsurance", "Cashflows")  # starts with the sheet name but is real
    assert not is_weak_name("Mortality C4 Rates", "Assumptions")


def test_assistant_names_are_used_only_where_the_heuristic_is_weak(tmp_path, section_heading_grids_xlsx):
    """A proposed name replaces a weak deterministic name, and is ignored when
    the surroundings already gave a real one."""
    from app.prep import plan_create_grid_titles

    result, _ = _run(tmp_path, section_heading_grids_xlsx)
    analysis, report = result["workbook_analysis"], result["validation_report"]
    # C3:D5 already gets a real name from its heading; B8:C8 only gets the
    # '<Sheet> <Anchor>' fallback, so only that one takes the proposal.
    baseline = {o["after"] for o in plan_create_grid_titles(analysis, report)[0] if o["op"] == "set_value"}
    assert "#CF B8" in baseline and "#Demographic Assumptions" in baseline

    proposed = {"CF!C3:D5": "Policy Demographics", "CF!B8:C8": "Section Two Heading"}
    names = {o["after"] for o in plan_create_grid_titles(analysis, report, proposed)[0] if o["op"] == "set_value"}
    assert "#Demographic Assumptions" in names  # heading wins over the proposal
    assert "#Policy Demographics" not in names
    assert "#Section Two Heading" in names      # weak fallback replaced by the proposal
    assert "#CF B8" not in names


def test_grid_naming_builds_context_and_parses_a_batched_answer(tmp_path, section_heading_grids_xlsx):
    from app.grid_naming import build_grid_context, suggest_names
    from app.grids import all_grids

    result, _ = _run(tmp_path, section_heading_grids_xlsx)
    analysis = result["workbook_analysis"]
    grid = next(g for g in all_grids(analysis) if g["ref"] == "C3:D5")
    ctx = build_grid_context(analysis, grid)
    assert ctx["id"] == "CF!C3:D5" and ctx["size"] == "3 rows x 2 cols"
    assert any("Demographic Assumptions" in line for line in ctx["above"])
    assert any("Gender" in line for line in ctx["sample"])

    calls: list[dict] = []

    def fake_completion(messages, system_prompt, **kw):
        calls.append({"messages": messages, "kw": kw})
        return {"available": True, "text": 'Here you go:\n{"CF!C3:D5": "Policy Demographics"}', "message": None}

    out = suggest_names([ctx], completion=fake_completion)
    assert out["available"] and out["names"] == {"CF!C3:D5": "Policy Demographics"}
    assert "CF!C3:D5" in calls[0]["messages"][0]["content"]


def test_grid_naming_reports_an_unavailable_gateway_without_raising():
    from app.grid_naming import suggest_names

    def unavailable(messages, system_prompt, **kw):
        return {"available": False, "text": None, "message": "no secret.key found nearby"}

    out = suggest_names([{"id": "S!A1", "size": "1 rows x 1 cols", "flags": [], "above": [], "left": [], "sample": ["x"]}], completion=unavailable)
    assert out["available"] is False and out["names"] == {} and "secret.key" in out["message"]
